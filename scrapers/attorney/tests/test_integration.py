import base64
import io
import threading

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from attorneyscraper.engine import Manager
from attorneyscraper.export import csv_bytes
from attorneyscraper.extract import extract
from attorneyscraper.models import RunConfig
from attorneyscraper.store import Store
from attorneyscraper.web import create_app

CONFIG = {
    "name": "Integration run",
    "sources": [{"organization": "Firm", "url": "https://firm.com/"}],
    "max_pages": 20,
    "delay": 0.5,
}


def wait(manager):
    manager.thread.join(timeout=10)
    assert not manager.thread.is_alive()


def test_end_to_end_crawl_review_export(tmp_path, factory):
    app = create_app(tmp_path / "test.db", factory)
    with TestClient(app) as c:
        token = c.get("/api/session").json()["token"]
        headers = {"x-session-token": token}
        response = c.post("/api/runs", json=CONFIG, headers=headers)
        assert response.status_code == 201, response.text
        rid = response.json()["id"]
        wait(app.state.manager)
        run = c.get("/api/runs/" + rid).json()
        assert run["status"] == "completed_with_issues"  # explicitly blocked /private
        data = c.get("/api/directory").json()
        assert data["total"] == 3
        assert all(p["status"] == "confirmed" for p in data["items"])
        jane = next(p for p in data["items"] if p["name"] == "Jane Smith")
        assert c.get("/api/directory?category=nurse").json()["total"] == 1
        review = c.post(
            "/api/review",
            params={"rid": rid, "key": jane["key"]},
            json={"status": "excluded", "note": "Manual audit: exclude this example"},
            headers=headers,
        )
        assert review.status_code == 200
        assert c.get("/api/directory?status=excluded").json()["total"] == 1
        x = c.get("/api/export/xlsx", params={"rid": rid})
        assert x.status_code == 200
        wb = load_workbook(io.BytesIO(x.content))
        assert wb.sheetnames == [
            "Summary",
            "Attorneys",
            "Evidence",
            "Pages",
            "Review history",
        ]
        assert wb["Attorneys"].max_row == 4
        assert wb["Evidence"].max_row > 4
        assert (
            c.get("/api/export/csv?status=excluded").content.count(b"Jane Smith") == 1
        )
        assert any(
            a["action"] == "review" for a in c.get("/api/runs/" + rid + "/audit").json()
        )
        assert c.get("/api/runs/missing").status_code == 404


def test_csrf_and_host_protection(tmp_path):
    with TestClient(create_app(tmp_path / "test.db")) as c:
        assert c.post("/api/runs", json=CONFIG).status_code == 403
        assert (
            c.get("/api/session", headers={"origin": "https://evil.com"}).status_code
            == 403
        )
        assert c.get("/api/session", headers={"host": "evil.com"}).status_code == 400


def test_upload_and_validation(tmp_path):
    with TestClient(create_app(tmp_path / "test.db")) as c:
        h = {"x-session-token": c.get("/api/session").json()["token"]}
        payload = {
            "filename": "sources.csv",
            "data": base64.b64encode(
                b"Organization,URL\nFirm,https://firm.com"
            ).decode(),
        }
        assert (
            c.post("/api/sources/validate", json=payload, headers=h).json()["sources"][
                0
            ]["organization"]
            == "Firm"
        )
        assert (
            c.post(
                "/api/runs", json={**CONFIG, "categories": ["unknown"]}, headers=h
            ).status_code
            == 422
        )


def test_resume_after_restart(tmp_path, factory):
    db = tmp_path / "test.db"
    store = Store(db)
    rid = store.create(RunConfig(**CONFIG))
    store.next_page(rid, "firm.com")
    store.set_status(rid, "running")
    restarted = Store(db)
    restarted.recover()
    assert restarted.get(rid)["status"] == "interrupted"
    assert restarted.pages(rid)[0]["status"] == "queued"
    manager = Manager(restarted, factory)
    manager.start(rid)
    wait(manager)
    assert len(restarted.records(rid)) == 3


def test_shared_phone_demoted_and_duplicate_profiles_merged(tmp_path):
    s = Store(tmp_path / "test.db")
    rid = s.create(RunConfig(**CONFIG))
    page = s.next_page(rid, "firm.com")
    records = []
    for name in ["Jane Smith", "John Jones"]:
        records += extract(
            f"<h1>{name}</h1><p>Attorney handling healthcare compliance.</p><p>Direct: 212-555-2345</p>",
            "https://firm.com/" + name.replace(" ", "-"),
            "Firm",
        )["records"]
    s.finish_page(page, {"records": records})
    assert all(
        c["owner"] == "organization" for p in s.records(rid) for c in p["contacts"]
    )
    s.finish_page(page, {"records": records})
    assert len(s.records(rid)) == 2


def test_csv_formula_injection():
    p = extract(
        "<h1>Jane Smith</h1><p>Attorney handles HIPAA</p>",
        "https://firm.com/jane",
        '=HYPERLINK("evil")',
    )["records"][0]
    assert b"'=HYPERLINK" in csv_bytes([p])


def test_pause_resume_and_concurrent_run_rejection(tmp_path, factory):
    ready = threading.Event()

    def slow(**kwargs):
        f = factory(**kwargs)
        old = f.get

        def get(url, site):
            ready.set()
            f.wait(0.3)
            return old(url, site)

        f.get = get
        return f

    s = Store(tmp_path / "test.db")
    rid = s.create(RunConfig(**CONFIG))
    m = Manager(s, slow)
    m.start(rid)
    assert ready.wait(2)
    import pytest

    with pytest.raises(ValueError):
        m.start(rid)
    m.control(rid, "pause")
    wait(m)
    assert s.get(rid)["status"] == "paused"
    assert s.pages(rid)[0]["status"] == "queued"
    m.fetcher_factory = factory
    m.start(rid)
    wait(m)
    assert len(s.records(rid)) == 3
