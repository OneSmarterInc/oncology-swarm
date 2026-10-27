import time

from attorneyscraper.engine import Manager
from attorneyscraper.extract import extract
from attorneyscraper.models import RunConfig, Source
from attorneyscraper.store import Store


def test_nested_label_fax_and_office_schema():
    html = """<script type="application/ld+json">{"@type":"LegalService","telephone":"212-869-3500","faxNumber":"212-398-1532"}</script>
    <h1>Eli Fuchsberg</h1><p>Attorney represents plaintiffs in medical malpractice.</p>
    <div class="item"><div>Phone</div><div>(212) 869-3500</div></div>
    <div class="item"><div>Fax</div><div>(212) 398-1532</div></div>
    <p>e.fuchsberg@fuchsberg.com</p>"""
    p = extract(html, "https://www.fuchsberg.com/attorneys/eli-fuchsberg")["records"][0]
    assert all(c["value"] != "+1 212-398-1532" for c in p["contacts"])
    assert (
        next(c for c in p["contacts"] if c["value"] == "+1 212-869-3500")["owner"]
        == "organization"
    )


def test_service_links_not_person_names():
    html = (
        "<h1>Eli Fuchsberg</h1><p>Attorney represents plaintiffs in medical malpractice.</p>"
        + "".join(
            f'<a href="/{i}">{name}</a>'
            for i, name in enumerate(
                [
                    "Personal Injury",
                    "Birth Injuries",
                    "Civil Rights",
                    "Product Liability",
                ]
            )
        )
    )
    assert (
        len(extract(html, "https://firm.com/attorneys/eli-fuchsberg")["records"]) == 1
    )


def test_h2_profile_and_same_name_different_profiles():
    p = extract(
        "<h2>Ilana Char</h2><p>Attorney handles healthcare transactions.</p>",
        "https://firm.com/people/ilana-char",
    )["records"][0]
    assert p["status"] == "confirmed"
    a = extract(
        "<h1>Jane Smith</h1><p>Attorney handles HIPAA</p>", "https://firm.com/jane-a"
    )["records"][0]
    b = extract(
        "<h1>Jane Smith</h1><p>Attorney handles HIPAA</p>", "https://firm.com/jane-b"
    )["records"][0]
    assert a["key"] != b["key"]


def test_wrong_practice_filter_excluded(tmp_path, factory):
    s = Store(tmp_path / "x.db")
    r = s.create(
        RunConfig(
            sources=[Source(organization="Firm", url="https://firm.com/people/jane")],
            max_depth=0,
            categories=["nurse"],
        )
    )
    m = Manager(s, factory)
    m.start(r)
    m.thread.join(5)
    assert s.records(r)[0]["status"] == "excluded"
    assert "Outside" in s.records(r)[0]["reason"]


def test_budget_reported_and_sources_first(tmp_path, factory):
    s = Store(tmp_path / "x.db")
    r = s.create(
        RunConfig(
            sources=[
                Source(organization="Firm", url="https://firm.com/"),
                Source(organization="Firm", url="https://firm.com/people/john"),
            ],
            max_pages=2,
        )
    )
    m = Manager(s, factory)
    m.start(r)
    m.thread.join(5)
    run = s.get(r)
    assert run["sites"][0]["reason"] == "page_limit"
    assert run["status"] == "completed_with_issues"
    assert s.records(r)[0]["name"] == "John Jones"


def test_parallel_sites_improve_elapsed_without_changing_results(tmp_path, factory):
    def run(workers, filename):
        def delayed(**kwargs):
            f = factory(**kwargs)
            old = f.get

            def get(url, host):
                f.wait(0.12)
                return old(url, host)

            f.get = get
            return f

        s = Store(tmp_path / filename)
        cfg = RunConfig(
            sources=[
                Source(organization=h, url=f"https://{h}/people/jane")
                for h in ["firm.com", "second.com", "third.com"]
            ],
            parallel_sites=workers,
            max_depth=0,
        )
        rid = s.create(cfg)
        m = Manager(s, delayed)
        began = time.monotonic()
        m.start(rid)
        m.thread.join(10)
        return time.monotonic() - began, len(s.records(rid))

    serial, count1 = run(1, "serial.db")
    parallel, count2 = run(3, "parallel.db")
    assert count1 == count2 == 3
    assert parallel < serial * 0.85


def test_html_challenge_is_not_person():
    assert (
        extract(
            "<h1>Access Denied</h1><p>Please enable JavaScript.</p>", "https://firm.com"
        )["records"]
        == []
    )


def test_source_whitespace_uses_domain():
    assert Source(organization="   ", url="https://firm.com").organization == "firm.com"


def test_named_hero_header_preserved_and_firm_news_removed():
    html = """<header><h1>Jane Smith</h1><p>Attorney</p></header><main><p>Jane handles healthcare transactions.</p>
    <section class="firm-news">Other attorneys won healthcare fraud and False Claims Act cases.</section></main>"""
    # Hero and main can be separate siblings, so name discovery includes the retained hero.
    records = extract(html, "https://firm.com/people/jane-smith")["records"]
    assert len(records) == 1
    assert records[0]["categories"] == ["transactional"]


def test_name_only_schema_widget_is_not_a_person_record():
    html = '<script type="application/ld+json">{"@type":"Person","name":"Alan Fuchsberg"}</script>'
    for path in ["contact", "services", "news"]:
        assert extract(html, f"https://firm.com/{path}")["records"] == []


def test_schema_person_uses_named_profile_link():
    html = """<h1>Our team</h1>
    <a href="/attorneys/alan-fuchsberg">Alan Fuchsberg</a>
    <script type="application/ld+json">{"@type":"Person","name":"Alan Fuchsberg"}</script>"""
    a = extract(html, "https://firm.com/team")["records"][0]
    b = extract(html, "https://firm.com/contact")["records"][0]
    assert a["profile_url"] == "https://firm.com/attorneys/alan-fuchsberg"
    assert a["key"] == b["key"]


def test_legacy_name_only_duplicates_hidden_but_reviews_preserved(tmp_path):
    import json

    store = Store(tmp_path / "legacy.db")
    rid = store.create(
        RunConfig(sources=[{"organization": "Firm", "url": "https://firm.com/"}])
    )
    p = extract("<h1>Jane Smith</h1><p>Attorney</p>", "https://firm.com/jane-smith")[
        "records"
    ][0]
    p.update(is_profile=False, attorney_evidence=False, categories=[], contacts=[])
    with store.db() as db:
        db.execute(
            "INSERT INTO observations VALUES(?,?,?,?)",
            (rid, p["source_url"], p["key"], json.dumps(p)),
        )
    assert store.records(rid) == []
    with store.db() as db:
        db.execute(
            "INSERT INTO reviews VALUES(?,?,?,?,?)",
            (rid, p["key"], "needs_review", "Keep for investigation", "2026-10-24"),
        )
    assert len(store.records(rid)) == 1


def test_glued_credentials_and_entity_suffixes_leave_the_name():
    from attorneyscraper.extract import clean_name

    assert clean_name("John Doe, P.C.") == "John Doe"  # was dropped entirely
    assert clean_name("Jane Smith PLLC") == "Jane Smith"
    assert clean_name("Jane Smith CPA, MBA") == "Jane Smith"
    assert clean_name("Mary Brown, LL.M.") == "Mary Brown"
    assert clean_name("Jane Smith, Jr.") == "Jane Smith Jr."  # generational suffix is part of the name
    assert clean_name("Mary Ann Pace") == "Mary Ann Pace"
    assert clean_name("Pam Pacheco") == "Pam Pacheco"


def test_excel_phone_cells_have_no_apostrophe(tmp_path):
    import io

    from openpyxl import load_workbook

    from attorneyscraper.export import csv_bytes, workbook_bytes

    record = {
        "name": "Jane Smith", "organization": "Firm", "title": "Attorney", "categories": [], "client_side": [],
        "contacts": [{"kind": "phone", "value": "+1 212-555-2345", "owner": "individual", "source_url": "u", "evidence": "Direct: 212-555-2345"}],
        "city": "", "state": "", "profile_url": "https://firm.com/people/jane", "observed_at": "",
        "evidence_urls": [], "evidence": [], "status": "confirmed", "is_profile": True, "key": "k",
        "attorney_evidence": [], "classification_evidence": {}, "automatic_status": "confirmed", "run_id": "r",
    }
    ws = load_workbook(io.BytesIO(workbook_bytes([record])))["Attorneys"]
    row = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]]))
    assert row["Individual phone"] == "+1 212-555-2345"
    assert "'+1 212-555-2345" in csv_bytes([record]).decode("utf-8-sig")  # CSV keeps the formula guard


def test_one_person_is_one_row_across_languages_and_the_directory_page(tmp_path):
    """fuchsberg.com returned Alan Fuchsberg four times: card, page, and both in Spanish."""
    from attorneyscraper.models import RunConfig, person_key, without_language
    from attorneyscraper.store import Store

    base = "https://www.fuchsberg.com"
    assert without_language(f"{base}/es/attorneys/alan-l-fuchsberg") == f"{base}/attorneys/alan-l-fuchsberg"
    assert person_key("Alan Fuchsberg", f"{base}/es/attorneys/alan-l-fuchsberg") == \
           person_key("Alan Fuchsberg", f"{base}/attorneys/alan-l-fuchsberg")
    # A real section is not mistaken for a language.
    assert without_language(f"{base}/practice-areas/malpractice") == f"{base}/practice-areas/malpractice"

    store = Store(tmp_path / "t.db")
    with store.owner:
        rid = store.create(RunConfig(name="dupes", sources=[{"organization": "Fuchsberg", "url": base + "/"}]))
        seen = [
            (f"{base}/attorneys", f"{base}/attorneys", False, "Managing Partner", []),
            (f"{base}/attorneys/alan-l-fuchsberg", f"{base}/attorneys/alan-l-fuchsberg", True, "Attorney",
             [{"kind": "email", "value": "alan@fuchsberg.com", "owner": "individual",
               "source_url": f"{base}/attorneys/alan-l-fuchsberg", "evidence": "Email alan@fuchsberg.com"}]),
            (f"{base}/es/attorneys", f"{base}/es/attorneys", False, "", []),
            (f"{base}/es/attorneys/alan-l-fuchsberg", f"{base}/es/attorneys/alan-l-fuchsberg", True, "", []),
        ]
        for source_url, profile_url, is_profile, title, contacts in seen:
            page = {"run_id": rid, "url": source_url, "host": "fuchsberg.com", "depth": 0, "id": None}
            with store.db() as db:
                db.execute("INSERT OR IGNORE INTO pages(run_id,host,url,depth,priority) VALUES(?,?,?,?,?)",
                           (rid, "fuchsberg.com", source_url, 0, 50))
                page["id"] = db.execute("SELECT id FROM pages WHERE run_id=? AND url=?",
                                        (rid, source_url)).fetchone()["id"]
            store.finish_page(page, {"records": [{
                "key": person_key("Alan Fuchsberg", profile_url), "name": "Alan Fuchsberg",
                "organization": "The Fuchsberg Law Firm", "title": title, "categories": [],
                "client_side": [], "contacts": contacts, "city": "New York", "state": "NY",
                "profile_url": profile_url, "source_url": source_url, "is_profile": is_profile,
                "status": "confirmed", "attorney_evidence": ["Attorney"], "classification_evidence": {},
                "observed_at": "2026-10-27T00:00:00+00:00",
            }]})
        people = [p for p in store.records(rid) if p["name"] == "Alan Fuchsberg"]
    assert len(people) == 1, [p["profile_url"] for p in people]
    kept = people[0]
    assert kept["profile_url"] == f"{base}/attorneys/alan-l-fuchsberg"   # the person's own page wins
    assert kept["title"] == "Attorney"
    assert any(c["value"] == "alan@fuchsberg.com" for c in kept["contacts"])
    assert f"{base}/attorneys" in kept["evidence_urls"]                  # the card is kept as evidence
