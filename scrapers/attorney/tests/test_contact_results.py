"""Contact results do not require healthcare classification approval."""

import csv
import io

import pytest
from openpyxl import load_workbook

from attorneyscraper.export import csv_bytes, workbook_bytes
from attorneyscraper.extract import extract


@pytest.mark.parametrize(
    "heading", ["Affordable Housing", "Fraud Alert", "Intellectual Property"]
)
def test_practice_and_alert_headings_are_not_people(heading):
    result = extract(
        f'<h1>{heading}</h1><p>Partner and counsel</p><a href="mailto:office@firm.com">Email</a>',
        "https://firm.com/services/example",
    )
    assert not result["records"]


def test_unclassified_person_contacts_export_without_review_labels():
    records = extract(
        '<h1>Barbara D. Linney</h1><p>Partner</p><a href="mailto:blinney@firm.com">Email</a>',
        "https://firm.com/people/barbara-linney",
    )["records"]
    assert records and records[0]["status"] == "needs_review"
    records[0].update(organization="Firm", city="", state="")
    text = csv_bytes(records).decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    assert "Classification" not in rows[0]
    assert "Insufficient evidence" not in text
    assert "blinney@firm.com" in text
    assert len(rows[0]) == len(rows[1])
    book = load_workbook(io.BytesIO(workbook_bytes(records)))
    assert "Classification" not in next(book["Attorneys"].values)
    assert "Insufficient evidence" not in str(list(book["Summary"].values))
    book.close()


def test_completed_export_keeps_specific_page_failure_details():
    pages = [
        {
            "url": "https://firm.com/large",
            "status": "error",
            "note": "FetchError: Page exceeds 3 MB limit",
            "record_count": 0,
            "elapsed": 1,
        }
    ]
    book = load_workbook(
        io.BytesIO(
            workbook_bytes([], run={"status": "completed_with_issues"}, pages=pages)
        )
    )
    assert dict(list(book["Summary"].values)[1:])["Run status"] == "Completed"
    assert list(book["Pages"].values)[1][2] == "FetchError: Page exceeds 3 MB limit"
    book.close()
