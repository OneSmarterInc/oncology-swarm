"""Portable application exports. Published contacts retain their ownership labels."""

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = [
    "Name",
    "Organization",
    "Title",
    "Categories",
    "Client side",
    "Review status",
    "Individual email",
    "Individual phone",
    "Organization email",
    "Organization phone",
    "Support contacts",
    "Unassigned contacts",
    "City",
    "State",
    "Profile URL",
    "Observed at",
    "Review note",
]


def safe(value):
    value = str(value if value is not None else "")
    value = "".join(c for c in value if ord(c) >= 32 or c in "\n\t")
    return "'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value


def rows(records):
    for p in records:

        def values(owner, kind=None):
            return "; ".join(
                sorted(
                    {
                        c["value"]
                        for c in p["contacts"]
                        if c["owner"] == owner and (not kind or c["kind"] == kind)
                    }
                )
            )

        yield [
            p["name"],
            p["organization"],
            p["title"],
            ", ".join(p["categories"]),
            ", ".join(p["client_side"]),
            p["status"],
            values("individual", "email"),
            values("individual", "phone"),
            values("organization", "email"),
            values("organization", "phone"),
            values("support_staff"),
            values("unassigned"),
            p["city"],
            p["state"],
            p["profile_url"],
            p["observed_at"],
            p.get("review", {}).get("note", ""),
        ]


def csv_bytes(records):
    out = io.StringIO(newline="")
    w = csv.writer(out)
    w.writerow(HEADERS)
    w.writerows([[safe(v) for v in row] for row in rows(records)])
    return out.getvalue().encode("utf-8-sig")


def workbook_bytes(records, run=None, pages=(), audit=()):
    book = Workbook()
    book.remove(book.active)

    def sheet(name, headers, data):
        ws = book.create_sheet(name)
        ws.append(headers)
        for row in data:
            ws.append([safe(v) if not isinstance(v, (int, float)) else v for v in row])
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="173B40")
            cell.font = Font(color="FFFFFF", bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = (
                40 if "URL" in header or "Evidence" in header else 25
            )
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        return ws

    sheet(
        "Summary",
        ["Metric", "Value"],
        [
            ("Run", (run or {}).get("name", "Directory")),
            ("Run status", (run or {}).get("status", "Snapshot")),
            ("People", len(records)),
            ("Confirmed", sum(p["status"] == "confirmed" for p in records)),
            ("Needs review", sum(p["status"] == "needs_review" for p in records)),
            (
                "Individual email",
                sum(
                    any(
                        c["owner"] == "individual" and c["kind"] == "email"
                        for c in p["contacts"]
                    )
                    for p in records
                ),
            ),
            (
                "Verification",
                "Published on source; mailbox deliverability and bar status are not verified",
            ),
            (
                "Coverage",
                "Completion refers to configured crawl, not every attorney on the website",
            ),
        ],
    )
    sheet("Attorneys", HEADERS, rows(records))
    evidence = []
    for p in records:
        for c in p["contacts"]:
            evidence.append(
                [
                    p["name"],
                    c["kind"],
                    c["value"],
                    c["owner"],
                    c["source_url"],
                    c["evidence"],
                    p["observed_at"],
                ]
            )
        for category, text in p["classification_evidence"].items():
            evidence.append(
                [
                    p["name"],
                    "category",
                    category,
                    "biography",
                    p["source_url"],
                    text,
                    p["observed_at"],
                ]
            )
    sheet(
        "Evidence",
        ["Name", "Field", "Value", "Owner", "Source URL", "Evidence", "Observed at"],
        evidence,
    )
    sheet(
        "Pages",
        ["URL", "Status", "Note", "People found", "Seconds"],
        [
            [p["url"], p["status"], p["note"], p["record_count"], p["elapsed"]]
            for p in pages
        ],
    )
    sheet(
        "Review history",
        ["Action", "Detail", "Time"],
        [[a["action"], a["detail"], a["created_at"]] for a in audit],
    )
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def parse_sources(filename, data):
    from .models import Source

    if len(data) > 2_000_000:
        raise ValueError("Source file must be smaller than 2 MB")
    if filename.lower().endswith(".csv"):
        reader = csv.DictReader(io.StringIO(data.decode("utf-8-sig")))
        raw = list(reader)
    elif filename.lower().endswith(".xlsx"):
        import zipfile

        from openpyxl import load_workbook

        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > 15_000_000:
                raise ValueError("Expanded workbook is too large")
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        iterator = wb.active.iter_rows(values_only=True)
        headers = next(iterator, [])
        raw = [dict(zip(headers, row)) for row in iterator]
        wb.close()
    else:
        raise ValueError("Upload a CSV or XLSX source list")
    result = []
    for number, row in enumerate(raw, 2):
        if not any(row.values()):
            continue
        row = {str(k).strip().lower(): str(v or "").strip() for k, v in row.items()}
        if row.get("enabled", "").lower() in {"no", "false", "0"}:
            continue
        org = row.get("organization") or row.get("firm name") or row.get("name", "")
        url = row.get("url") or row.get("website") or row.get("official website", "")
        try:
            result.append(
                Source(
                    organization=org,
                    url=url,
                    city=row.get("city", ""),
                    state=row.get("state", ""),
                )
            )
        except ValueError as e:
            raise ValueError(
                f"Row {number}: organization and a valid URL are required"
            ) from e
    if not result or len(result) > 100:
        raise ValueError("Provide 1–100 enabled source rows")
    return result
