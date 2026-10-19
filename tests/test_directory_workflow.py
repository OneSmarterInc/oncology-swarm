from copy import deepcopy
from dataclasses import replace, asdict
from datetime import timedelta
from io import BytesIO
import pytest
from openpyxl import load_workbook
from sqlalchemy import select, func
from app.cli import SEEDS
from app.worker import Coordinator
from app.service import Service,Conflict
from app.models import Job,Proposal,Snapshot,Doctor,DoctorSource,Control,Audit,LLMUsage,HostPacing
from app.fetcher import SafeFetcher
from app.extractor import extract
from app.export import export_xlsx

def proposed(stack):
    settings,db,service=stack
    ids=[service.enqueue(url)[0] for url in SEEDS]
    Coordinator(service).run(once=True)
    return ids

def approved(stack):
    ids=proposed(stack)
    for id in ids:stack[2].review(id,1,'Reviewer','Checked each source',True)
    return ids

def test_correction_withdrawal_and_audit(stack):
    _,db,service=stack;ids=approved(stack)
    assert load_workbook(BytesIO(export_xlsx(db)))['Doctors'].max_row==3
    assert service.correct(ids[0],2,'Reviewer','Recheck identity')=='review'
    assert load_workbook(BytesIO(export_xlsx(db)))['Doctors'].max_row==2
    with db.session() as s:p=s.scalar(select(Proposal).where(Proposal.job_id==ids[0]));data=deepcopy(p.data)
    assert not service.edit(ids[0],data,3)
    service.review(ids[0],4,'Reviewer','Rechecked',True)
    assert load_workbook(BytesIO(export_xlsx(db)))['Doctors'].max_row==3
    service.withdraw(ids[0],5,'Reviewer','Withdraw at request')
    assert load_workbook(BytesIO(export_xlsx(db)))['Doctors'].max_row==2
    with db.session() as s:
        assert s.scalar(select(func.count()).select_from(Doctor))==2
        assert s.scalar(select(Audit).where(Audit.action=='correct')).details['previous_data']==data
    with pytest.raises(Conflict):service.correct(ids[0],5,'Reviewer','Stale')

def test_refresh_unchanged_keeps_original_without_new_proposal(stack):
    _,db,service=stack;id=approved(stack)[0]
    result=service.refresh_job(id,2,'Reviewer','Freshness check')
    assert result['id']!=id and result['generation']==1
    Coordinator(service).run(once=True)
    with db.session() as s:
        assert s.get(Job,result['id']).state=='unchanged'
        assert s.scalar(select(Snapshot).where(Snapshot.job_id==result['id']))
        assert not s.scalar(select(Proposal).where(Proposal.job_id==result['id']))
        assert s.get(Job,id).state=='approved'
    assert load_workbook(BytesIO(export_xlsx(db)))['Doctors'].max_row==3

def test_changed_refresh_quarantines_old_until_review(stack):
    settings,db,service=stack;id=approved(stack)[0]
    fresh=service.refresh_job(id,2,'Reviewer','Freshness check')['id']
    service.claim('refresh')
    fetched=SafeFetcher(settings,service.reserve_request).fetch(SEEDS[0])
    fetched=replace(fetched,content_hash='f'*64)
    assert service.finish(fresh,'refresh',fetched,extract(fetched))
    assert load_workbook(BytesIO(export_xlsx(db)))['Doctors'].max_row==2
    with db.session() as s:assert not s.scalar(select(Proposal).where(Proposal.job_id==fresh)).duplicate_ids
    service.review(fresh,1,'Reviewer','Changed source checked',True)
    assert load_workbook(BytesIO(export_xlsx(db)))['Doctors'].max_row==3
    with db.session() as s:
        assert s.get(Job,id).state=='superseded'
        assert s.get(Job,fresh).state=='approved'

def test_refresh_cannot_duplicate_pending_work(stack):
    _,db,service=stack;id=approved(stack)[0]
    service.refresh_job(id,2,'Reviewer','Refresh')
    with pytest.raises(Conflict,match='unfinished'):service.refresh_job(id,2,'Reviewer','Again')

def test_merge_retains_sources_and_withdraw_quarantines_group(stack):
    _,db,service=stack;ids=proposed(stack)
    with db.session.begin() as s:
        a=s.scalar(select(Proposal).where(Proposal.job_id==ids[0]));b=s.scalar(select(Proposal).where(Proposal.job_id==ids[1]))
        sa=s.scalar(select(Snapshot).where(Snapshot.job_id==ids[0]));sb=s.scalar(select(Snapshot).where(Snapshot.job_id==ids[1]))
        data=deepcopy(a.data);sb.text=sa.text
        for e in data['evidence']:e['source_url'],e['retrieved_at']=sb.url,sb.retrieved_at
        b.data=data
    for id in ids:service.review(id,1,'Reviewer','Same individual sources',True,True)
    merged=service.merge(ids[1],ids[0],2,2,'Reviewer','Matched identity across sources')
    wb=load_workbook(BytesIO(export_xlsx(db)))
    assert wb['Doctors'].max_row==2 and wb['Sources'].max_row==3
    with db.session() as s:assert s.scalar(select(func.count()).select_from(DoctorSource).where(DoctorSource.doctor_id==merged['doctor_id']))==2
    with pytest.raises(Conflict,match='Merged'):service.correct(ids[0],3,'Reviewer','Unsafe merged correction')
    service.withdraw(ids[1],3,'Reviewer','One source requested correction')
    assert load_workbook(BytesIO(export_xlsx(db)))['Doctors'].max_row==1

def test_merge_different_identities_blocked(stack):
    _,db,service=stack;ids=approved(stack)
    with pytest.raises(Conflict,match='Names differ'):service.merge(ids[1],ids[0],2,2,'Reviewer','No')

def test_run_budgets_reset_only_when_drained_and_paused(stack):
    _,db,service=stack
    id=service.enqueue(SEEDS[0])[0]
    with pytest.raises(Conflict):service.reset_budgets('Reviewer','New run')
    with db.session.begin() as s:s.get(Control,1).paused=True
    with pytest.raises(Conflict):service.reset_budgets('Reviewer','Pending work remains')
    with db.session.begin() as s:s.get(Control,1).paused=False
    Coordinator(service).run(once=True)
    with db.session.begin() as s:s.get(Control,1).paused=True
    result=service.reset_budgets('Reviewer','Approved next run')
    assert result['run_number']==2 and result['previous']['jobs_created']==1
    with db.session() as s:assert s.get(Control,1).jobs_created==0

def test_invalid_model_shapes_remain_reviewable(stack):
    settings,db,service=stack
    id=service.enqueue(SEEDS[0])[0];service.claim('test')
    fetched=SafeFetcher(settings,service.reserve_request).fetch(SEEDS[0])
    assert service.finish(id,'test',fetched,{'doctor':'bad','unknown':123})
    with db.session() as s:
        p=s.scalar(select(Proposal).where(Proposal.job_id==id));assert p.validation_errors and p.normalized_name==''
    with pytest.raises(Conflict,match='Validation'):service.review(id,1,'Reviewer','Cannot approve',True)

def test_usage_captured_per_call_even_manual_retries(stack):
    _,db,service=stack;id=service.enqueue(SEEDS[0])[0]
    usage={'provider':'mistral','model':'test','input_tokens':100,'output_tokens':20,'total_tokens':120,'estimated_cost_usd':0.001}
    for i in range(2):
        service.claim('test');service.reserve_llm_call(id,'test')
        assert service.record_llm_usage(id,'test',usage)
        assert not service.record_llm_usage(id,'test',usage)
        service.fail(id,'test','Invalid provider response')
        if i==0:service.retry_failed(id,True)
    metrics=service.metrics()
    assert metrics['llm_calls_reserved']==2 and metrics['llm_responses_recorded']==2
    assert metrics['input_tokens']==200 and metrics['estimated_cost_usd']==0.002
    assert not metrics['cost_is_partial']

def test_shared_pacing_respects_new_robots_delay(stack,monkeypatch):
    _,db,service=stack
    import app.service as module
    now=[module.utcnow()]
    monkeypatch.setattr(module,'utcnow',lambda:now[0])
    waits=[]
    def sleep(wait):
        waits.append(wait);now[0]+=timedelta(seconds=wait)
    monkeypatch.setattr('time.sleep',sleep)
    service.pace_host('hospital.example',10)
    other=Service(db,service.settings)
    other.pace_host('hospital.example',60)
    service.pace_host('hospital.example',10)
    assert waits==[60,60]
    with db.session() as s:assert s.get(HostPacing,'hospital.example').interval_seconds==60

def test_live_sources_are_approved_not_implicitly_trusted(stack):
    settings,db,_=stack;service=Service(db,replace(settings,mode='live',allowed_domains=('hospital.example',)))
    with pytest.raises(ValueError,match='approved'):service.enqueue('https://hospital.example/doctor')
    source=service.add_source('Example Hospital','Delhi','hospital.example','https://hospital.example/doctors','Reviewer','Official source checked')
    id,created=service.enqueue('https://hospital.example/doctor')
    assert created
    with db.session() as s:assert s.get(Job,id).source_id==source['id']
    with pytest.raises(ValueError):service.add_source('Invalid','Delhi','hospital.example','https://other.example/doctors','Reviewer')
    with pytest.raises(ValueError):service.add_source('Invalid','Delhi','hospital.example','https://hospital.example/other','Reviewer',profile_url_pattern='.*')

def test_discovery_lease_prevents_budget_reset_and_fences_stale_calls(stack):
    settings,db,_=stack
    from app.models import Source,Hospital,DiscoveryRun,utcnow
    service=Service(db,replace(settings,mode='live',allowed_domains=('hospital.example',)))
    source=service.add_source('Official','City','hospital.example','https://hospital.example/doctors','Reviewer')
    with db.session.begin() as s:
        run=DiscoveryRun(source_id=source['id'],state='running',owner='one',lease_expires_at=utcnow()+timedelta(seconds=60))
        s.add(run);s.flush();id=run.id;s.get(Control,1).paused=True
    with pytest.raises(Conflict,match='discovery'):service.reset_budgets('Reviewer','Next')
    with db.session.begin() as s:s.get(DiscoveryRun,id).lease_expires_at=utcnow()-timedelta(seconds=1)
    service.reset_budgets('Reviewer','Expired discovery recovered')
    with db.session.begin() as s:s.get(Control,1).paused=False
    with pytest.raises(Conflict,match='lease'):service.reserve_request('https://hospital.example/a',id,'one')
    with pytest.raises(Conflict,match='lease'):service.enqueue('https://hospital.example/a',source['id'],id,'one')

def test_refresh_approval_cannot_undo_withdrawal(stack):
    settings,db,service=stack;id=approved(stack)[0]
    fresh=service.refresh_job(id,2,'Reviewer','Check')['id'];service.claim('refresh')
    fetched=replace(SafeFetcher(settings,service.reserve_request).fetch(SEEDS[0]),content_hash='9'*64)
    service.finish(fresh,'refresh',fetched,extract(fetched))
    service.withdraw(id,2,'Reviewer','Remove record')
    with pytest.raises(Conflict,match='withdrawn'):service.review(fresh,1,'Reviewer','Cannot restore silently',True)

def test_changed_refresh_quarantines_before_model_failure(stack,monkeypatch):
    settings,db,service=stack;id=approved(stack)[0]
    fresh=service.refresh_job(id,2,'Reviewer','Check changed source')['id']
    result=replace(SafeFetcher(settings,service.reserve_request).fetch(SEEDS[0]),content_hash='4'*64)
    coordinator=Coordinator(service)
    monkeypatch.setattr(coordinator.fetcher,'fetch',lambda *a,**kw:result)
    class BrokenModel:
        def extract_with_usage(self,result):raise RuntimeError('Mocked provider failure')
    coordinator.llm=BrokenModel()
    coordinator.run(once=True)
    with db.session() as s:
        assert s.get(Job,fresh).state=='failed'
        assert s.scalar(select(Snapshot).where(Snapshot.job_id==fresh)).content_hash=='4'*64
        p=s.scalar(select(Proposal).where(Proposal.job_id==id))
        assert s.scalar(select(Doctor).where(Doctor.proposal_id==p.id)).needs_review
    assert load_workbook(BytesIO(export_xlsx(db)))['Doctors'].max_row==2
    with pytest.raises(Conflict,match='new refresh'):service.retry_failed(fresh,True)
    assert service.refresh_job(id,2,'Reviewer','Retry with retained prior evidence')['id']>fresh
