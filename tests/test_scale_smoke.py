"""Synthetic, offline batch smoke: never a real coverage/throughput claim."""
from dataclasses import replace
from datetime import datetime,timezone
import hashlib
from io import BytesIO
from openpyxl import load_workbook
from sqlalchemy import select,func,text
from app.service import Service
from app.models import Job,Proposal,Control,Doctor
from app.fetcher import FetchResult,visible_text
from app.worker import Coordinator
from app.export import export_xlsx

def test_two_hundred_synthetic_profiles_through_review_and_export(stack,monkeypatch):
    settings,db,_=stack
    settings=replace(settings,mode='live',allowed_domains=('synthetic.example',),max_jobs=220,max_requests=220)
    service=Service(db,settings)
    source=service.add_source('Synthetic Hospital','Test City','synthetic.example','https://synthetic.example/doctors','Synthetic test reviewer','Fictional test only; no network')
    ids=[]
    for i in range(200):ids.append(service.enqueue(f'https://synthetic.example/doctor/oncology-{i}',source['id'])[0])
    coordinator=Coordinator(service)
    def fetch(url,**kwargs):
        service.reserve_request(url)
        index=url.rsplit('-',1)[1]
        html=f'<p>Doctor name: Dr Synthetic Person {index}</p><p>Specialty: Medical Oncology</p><p>Hospital: Synthetic Hospital</p><p>City: Test City</p>'
        return FetchResult(url,visible_text(html),html,datetime.now(timezone.utc).isoformat(),hashlib.sha256(html.encode()).hexdigest(),True)
    monkeypatch.setattr(coordinator.fetcher,'fetch',fetch)
    coordinator.run(once=True)
    with db.session() as s:
        proposals=s.scalars(select(Proposal).order_by(Proposal.job_id)).all()
        assert len(proposals)==200
        assert all(not p.validation_errors and not p.duplicate_ids for p in proposals)
        assert s.scalar(select(func.count()).select_from(Job).where(Job.state=='review'))==200
        assert s.get(Control,1).requests_used==200
        plan=s.execute(text("EXPLAIN QUERY PLAN SELECT job_id FROM proposals WHERE normalized_name='synthetic person 100'")).all()
        assert any('ix_proposals_normalized_name' in str(row) for row in plan)
    # Programmatic test approvals are synthetic regression operations, not measured
    # human review throughput and not permission to approve real medical records.
    for start in range(0,200,25):
        result=service.bulk_approve([{'id':id,'expected_version':1} for id in ids[start:start+25]],'Synthetic test reviewer','Fictional evidence fixture checked',True)
        assert all(x.get('state')=='approved' for x in result)
    wb=load_workbook(BytesIO(export_xlsx(db)))
    assert wb['Doctors'].max_row==201 and wb['Sources'].max_row==201
    synthetic_column=list(wb['Doctors'].values)[0].index('synthetic')
    assert all(row[synthetic_column] is True for row in list(wb['Doctors'].values)[1:])
