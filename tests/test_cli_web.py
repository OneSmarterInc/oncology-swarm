import json
import time

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from oncoscraper import web
from oncoscraper.cli import main
from oncoscraper.extract import extract_page
from oncoscraper.models import Page


def test_cli_scrape_writes_excel(site_url, tmp_path, capsys, monkeypatch):
    out = tmp_path / "onc.xlsx"
    monkeypatch.setenv("ONCOSCRAPER_ALLOW_PRIVATE", "1")
    rc = main(["scrape", site_url, "--out", str(out), "--delay", "0", "--render", "never", "--max-pages", "60"])
    assert rc == 0
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Oncologists", "All doctors", "Evidence", "Pages"]
    names = [r[0].value for r in wb["Oncologists"].iter_rows(min_row=2)]
    assert "Dr. Meenu Walia" in names and "Dr. Sunil Kumar" not in names
    assert wb["Evidence"].max_row > 50
    assert "oncologists" in capsys.readouterr().out


def test_cli_file_command(tmp_path, capsys):
    p = tmp_path / "p.html"
    p.write_text("<!-- source: https://h.example/doctor/dr-x-y -->\n<html><body><h1>Dr. Xavier Yadav</h1>"
                 "<p>Consultant - Radiation Oncology</p><p>MBBS, MD (Radiotherapy)</p></body></html>")
    assert main(["file", str(p)]) == 0
    out = capsys.readouterr().out
    assert "Dr. Xavier Yadav" in out and "Radiation Oncology" in out


def test_web_ui_job_flow(site_url, tmp_path, monkeypatch):
    monkeypatch.setattr(web, "OUT_DIR", tmp_path)
    monkeypatch.setenv("ONCOSCRAPER_ALLOW_PRIVATE", "1")
    client = TestClient(web.app)
    assert "OncoScraper" in client.get("/").text
    jid = client.post("/api/jobs", json={"urls": [site_url], "delay_seconds": 0, "render": "never", "max_pages": 60}).json()["id"]
    for _ in range(100):
        st = client.get(f"/api/jobs/{jid}").json()
        if st["state"] != "running":
            break
        time.sleep(0.1)
    assert st["state"] == "done", st
    assert st["progress"]["oncologists"] >= 11
    assert all(d["is_oncology"] for d in st["doctors"])
    r = client.get(f"/api/jobs/{jid}/download.xlsx")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert "Dr. Asha Rao" in client.get(f"/api/jobs/{jid}/download.csv").text
    assert json.loads(client.get(f"/api/jobs/{jid}/download.json").text)["stats"]["oncologists"] >= 11


def test_next_app_router_stream():
    payload = json.dumps({"doctor": {"name": "Dr. Omkar Pai", "speciality": "Medical Oncology", "designation": "Consultant",
                                     "qualification": "MBBS, MD, DM", "experience": "9 Years"}})
    chunk = json.dumps("5:" + payload)[1:-1]
    html = f'<html><body><div id="__next"></div><script>self.__next_f.push([1,"{chunk}"])</script></body></html>'
    pr = extract_page(Page("https://h.example/doctors/dr-omkar-pai", "https://h.example/doctors/dr-omkar-pai", 200, html))
    d = pr.doctors[0]
    assert d.get("name") == "Dr. Omkar Pai" and d.get("experience_years") == "9" and d.is_oncology


def test_jsonld_itemlist_listing():
    items = [{"@type": "Physician", "name": f"Dr. {n}", "medicalSpecialty": "Surgical Oncology",
              "url": f"/doctors/dr-{n.lower().replace(' ', '-')}"} for n in ("Ira Sen", "Om Rao", "Tara Das")]
    html = ('<html><head><script type="application/ld+json">' + json.dumps({"@type": "ItemList", "itemListElement": items})
            + '</script></head><body><h1>Our Oncologists</h1></body></html>')
    pr = extract_page(Page("https://h.example/oncologists", "https://h.example/oncologists", 200, html))
    assert pr.page_type == "listing" and len(pr.doctors) == 3
    assert {d.profile_url for d in pr.doctors} == {"https://h.example/doctors/dr-ira-sen", "https://h.example/doctors/dr-om-rao", "https://h.example/doctors/dr-tara-das"}


def test_batch_csv_with_resume(site_url, tmp_path, capsys, monkeypatch):
    monkeypatch.setenv("ONCOSCRAPER_ALLOW_PRIVATE", "1")
    csv_path = tmp_path / "hospitals.csv"
    csv_path.write_text(
        "hospital,official_domain,doctors_list_url,oncology_url\n"
        f"Sunrise Cancer Care,,,{site_url}/specialities/cancer-care\n"
        f"Sunrise Faculty,,{site_url}/faculty/radiation-oncology,\n"
        "No URL Hospital,,,\n", encoding="utf-8")
    out = tmp_path / "india.xlsx"
    args = ["batch", str(csv_path), "--out", str(out), "--delay", "0", "--render", "never", "--pages-per-site", "60", "--no-sitemap"]
    assert main(args) == 0
    wb = load_workbook(out)
    assert wb.sheetnames[:2] == ["Summary", "Hospitals"]
    hosp = {r[1].value: r for r in wb["Hospitals"].iter_rows(min_row=2)}
    assert set(hosp) == {"Sunrise Cancer Care", "Sunrise Faculty"}
    header = [c.value for c in wb["Oncologists"][1]]
    rows = [dict(zip(header, [c.value for c in r])) for r in wb["Oncologists"].iter_rows(min_row=2)]
    hc = next(r for r in rows if r["Name"] == "Dr. Harish Chandra" and r["Source Hospital"] == "Sunrise Faculty")
    assert hc["Doctor Email"] == "hchandra@inst.example" and hc["Source Hospital"] == "Sunrise Faculty"
    # Second run resumes from saved per-site results without crawling again.
    assert main(args) == 0
    assert "resumed" in capsys.readouterr().out


def test_excel_template_round_trip(tmp_path):
    from oncoscraper.sources import read_sources, write_template
    rows = [
        {"hospital_name": "Alpha Cancer Centre", "oncology_page_urls": "https://alpha.example/oncology\nhttps://alpha.example/team",
         "sample_profile_url": "https://alpha.example/doctor/dr-a-b", "max_pages": "40", "use_browser": "always"},
        {"hospital_name": "Beta Hospital", "official_website": "beta.example"},
        {"hospital_name": "Gamma", "doctors_list_urls": "https://gamma.example/list.pdf"},
        {"hospital_name": "Delta", "official_website": "https://delta.example", "enabled": "no"},
        {"official_website": "https://www.epsilon.example/", "sample_profile_url": "https://www.epsilon.example/dr-ravi-kant/"},
    ]
    path = write_template(tmp_path / "list.xlsx", rows)
    sources, problems = read_sources(path)
    by = {s.hospital: s for s in sources}
    a = by["Alpha Cancer Centre"]
    assert a.urls == ["https://alpha.example/oncology", "https://alpha.example/team", "https://alpha.example/doctor/dr-a-b"]
    assert a.patterns == ["/doctor/"] and a.max_pages == 40 and a.render == "always"
    assert by["Beta Hospital"].urls == ["https://beta.example/"]
    assert by["www.epsilon.example"].patterns == ["/dr-"]
    probs = {p.hospital: p.problem for p in problems}
    assert "PDF" not in probs.get("Gamma", "") and "no usable URL" in probs["Gamma"]
    assert probs["Delta"].startswith("skipped")


def test_web_batch_upload(site_url, tmp_path, monkeypatch):
    import base64
    from oncoscraper.sources import write_template
    monkeypatch.setattr(web, "OUT_DIR", tmp_path)
    monkeypatch.setenv("ONCOSCRAPER_ALLOW_PRIVATE", "1")
    path = write_template(tmp_path / "up.xlsx", [
        {"hospital_name": "Sunrise Faculty", "oncology_page_urls": f"{site_url}/faculty/radiation-oncology"},
        {"hospital_name": "Broken row"}])
    client = TestClient(web.app)
    assert client.get("/api/template.xlsx").content[:2] == b"PK"
    payload = {"filename": "up.xlsx", "content_base64": base64.b64encode(path.read_bytes()).decode()}
    check = client.post("/api/sources/check", json=payload).json()
    assert check["usable"] == 1 and check["problems"][0]["hospital"] == "Broken row"
    jid = client.post("/api/batch", json={**payload, "pages_per_site": 1, "delay_seconds": 0, "render": "never",
                                          "use_sitemaps": False}).json()["id"]
    for _ in range(100):
        st = client.get(f"/api/batch/{jid}").json()
        if st["state"] != "running":
            break
        time.sleep(0.1)
    assert st["state"] == "done", st
    assert st["sites"][0]["oncologists"] == 3 and st["sites"][0]["with_email"] == 2
    assert client.get(f"/api/batch/{jid}/download.xlsx").content[:2] == b"PK"


def test_partial_excel_while_batch_runs(site_url, tmp_path, monkeypatch):
    """The combined Excel can be downloaded at any time; it holds finished hospitals and the one in progress."""
    from oncoscraper.batch import run_batch, snapshot_xlsx
    from oncoscraper.config import ScrapeConfig
    from oncoscraper.sources import Source
    monkeypatch.setenv("ONCOSCRAPER_ALLOW_PRIVATE", "1")
    srcs = [Source(row=2, hospital="Sunrise Faculty", urls=[f"{site_url}/faculty/radiation-oncology"]),
            Source(row=3, hospital="Sunrise Cancer Care", urls=[f"{site_url}/specialities/cancer-care"])]
    live, snaps = {}, []

    def progress(e):
        if e["event"] == "page" and 3 in (live.get("active") or {}) and not snaps:
            snaps.append(snapshot_xlsx(live, tmp_path / "part.xlsx"))

    run_batch(srcs, tmp_path / "all.xlsx", ScrapeConfig(delay_seconds=0, render="never", use_sitemaps=False,
                                                        allow_private_hosts=True),
              per_site_pages=3, resume=False, progress=progress, live=live)
    wb = load_workbook(snaps[0])
    hosp = {r[1].value: r[2].value for r in wb["Hospitals"].iter_rows(min_row=2)}
    assert hosp == {"Sunrise Faculty": "partial", "Sunrise Cancer Care": "in progress"}
    names = {r[0].value for r in wb["Oncologists"].iter_rows(min_row=2)}
    assert "Dr. Harish Chandra" in names and len(names) > 3


def test_dead_links_are_reported_separately():
    from oncoscraper.batch import _failure_note
    pages = [{"note": "HTTP 404"}] * 9 + [{"note": "HTTP 500"}]
    note = _failure_note(pages, 150)
    assert note.startswith("1 of 150 pages could not be read") and "9 dead links on the site" in note
    assert _failure_note([{"note": "HTTP 404"}], 150) == "1 dead link on the site (404, page removed)"


def test_evidence_has_one_row_per_doctor_and_field(site_url, tmp_path, monkeypatch):
    from oncoscraper.batch import run_batch
    from oncoscraper.config import ScrapeConfig
    from oncoscraper.sources import Source
    out = run_batch([Source(row=2, hospital="Sunrise", urls=[f"{site_url}/specialities/cancer-care"])],
                    tmp_path / "e.xlsx", ScrapeConfig(delay_seconds=0, render="never", use_sitemaps=False,
                                                      allow_private_hosts=True), per_site_pages=15, resume=False)
    wb = load_workbook(out)
    pairs = [(r[0].value, r[1].value) for r in wb["Evidence"].iter_rows(min_row=2)]
    assert pairs and len(pairs) == len(set(pairs))
    names = [r[0].value for r in wb["Oncologists"].iter_rows(min_row=2)]
    assert len(names) == len(set(names))


def test_single_crawl_downloads_while_running(site_url, tmp_path, monkeypatch):
    monkeypatch.setattr(web, "OUT_DIR", tmp_path)
    monkeypatch.setenv("ONCOSCRAPER_ALLOW_PRIVATE", "1")
    client = TestClient(web.app)
    jid = client.post("/api/jobs", json={"urls": [site_url + "/specialities/cancer-care"], "max_pages": 40,
                                         "delay_seconds": 0.4, "render": "never", "use_sitemaps": False}).json()["id"]
    for _ in range(100):
        st = client.get(f"/api/jobs/{jid}").json()
        if st["progress"]["oncologists"] >= 2 or st["state"] != "running":
            break
        time.sleep(0.1)
    assert st["state"] == "running"
    x = client.get(f"/api/jobs/{jid}/download.xlsx")
    c = client.get(f"/api/jobs/{jid}/download.csv")
    j = client.get(f"/api/jobs/{jid}/download.json")
    assert x.status_code == c.status_code == j.status_code == 200
    assert x.content[:2] == b"PK" and "partial" in x.headers["content-disposition"]
    assert c.text.count("\n") >= 3 and '"partial"' in j.text
    client.post(f"/api/jobs/{jid}/stop")


def _wait(client, url, cond, tries=150):
    for _ in range(tries):
        st = client.get(url).json()
        if cond(st):
            return st
        time.sleep(0.1)
    return st


def test_pause_and_resume_single_crawl_and_batch(site_url, tmp_path, monkeypatch):
    import base64
    from oncoscraper.sources import write_template
    monkeypatch.setattr(web, "OUT_DIR", tmp_path)
    monkeypatch.setenv("ONCOSCRAPER_ALLOW_PRIVATE", "1")
    client = TestClient(web.app)
    # single crawl
    jid = client.post("/api/jobs", json={"urls": [site_url + "/specialities/cancer-care"], "max_pages": 30,
                                         "delay_seconds": 0.3, "render": "never", "use_sitemaps": False}).json()["id"]
    _wait(client, f"/api/jobs/{jid}", lambda s: s["progress"]["fetched"] >= 2)
    assert client.post(f"/api/jobs/{jid}/pause").json()["state"] == "paused"
    time.sleep(0.6)
    a = client.get(f"/api/jobs/{jid}").json()
    time.sleep(1.0)
    b = client.get(f"/api/jobs/{jid}").json()
    assert a["state"] == b["state"] == "paused" and a["progress"]["fetched"] == b["progress"]["fetched"]
    assert client.get(f"/api/jobs/{jid}/download.csv").status_code == 200  # download while paused
    client.post(f"/api/jobs/{jid}/resume")
    c = _wait(client, f"/api/jobs/{jid}", lambda s: s["progress"]["fetched"] > b["progress"]["fetched"])
    assert c["progress"]["fetched"] > b["progress"]["fetched"]
    client.post(f"/api/jobs/{jid}/stop")
    # batch
    path = write_template(tmp_path / "b.xlsx", [
        {"hospital_name": "Sunrise A", "oncology_page_urls": f"{site_url}/specialities/cancer-care"},
        {"hospital_name": "Sunrise B", "oncology_page_urls": f"{site_url}/faculty/radiation-oncology"}])
    payload = {"filename": "b.xlsx", "content_base64": base64.b64encode(path.read_bytes()).decode(),
               "pages_per_site": 4, "delay_seconds": 0.3, "render": "never", "use_sitemaps": False}
    bid = client.post("/api/batch", json=payload).json()["id"]
    _wait(client, f"/api/batch/{bid}", lambda s: s["progress"].get("fetched", 0) >= 1)
    client.post(f"/api/batch/{bid}/pause")
    time.sleep(0.6)
    p1 = client.get(f"/api/batch/{bid}").json()
    time.sleep(1.0)
    p2 = client.get(f"/api/batch/{bid}").json()
    assert p1["state"] == "paused" and p1["progress"] == p2["progress"] and p1["sites"] == p2["sites"]
    client.post(f"/api/batch/{bid}/resume")
    done = _wait(client, f"/api/batch/{bid}", lambda s: s["state"] == "done", tries=300)
    assert done["state"] == "done" and len(done["sites"]) == 2


def test_ui_helpers_and_friendly_errors(site_url, tmp_path, monkeypatch, capsys):
    import base64
    client = TestClient(web.app)
    assert "app.js" in client.get("/").text and "attachBatch" in client.get("/static/app.js").text
    assert set(client.get("/api/active").json()) == {"job", "batch"}
    bad = {"filename": "x.csv", "content_base64": base64.b64encode(b"foo,bar\n1,2\n").decode()}
    r = client.post("/api/sources/check", json=bad)
    assert r.status_code == 400 and "no URL column" in r.json()["detail"]
    p = tmp_path / "x.csv"
    p.write_text("foo,bar\n1,2\n")
    assert main(["batch", str(p), "--out", str(tmp_path / "o.xlsx")]) == 2
    assert "Cannot read" in capsys.readouterr().out


def test_portal_history_directory_dashboard_and_resume(site_url, tmp_path, monkeypatch):
    import base64
    from oncoscraper.sources import write_template
    monkeypatch.setattr(web, "OUT_DIR", tmp_path)
    monkeypatch.setenv("ONCOSCRAPER_ALLOW_PRIVATE", "1")
    monkeypatch.delenv("ONCOSCRAPER_DB", raising=False)
    web._STORE = None
    client = TestClient(web.app)
    path = write_template(tmp_path / "p.xlsx", [
        {"hospital_name": "Sunrise A", "oncology_page_urls": f"{site_url}/specialities/cancer-care"},
        {"hospital_name": "Sunrise B", "oncology_page_urls": f"{site_url}/faculty/radiation-oncology"}])
    payload = {"filename": "p.xlsx", "content_base64": base64.b64encode(path.read_bytes()).decode(),
               "pages_per_site": 60, "delay_seconds": 0, "render": "never", "use_sitemaps": False}
    bid = client.post("/api/batch", json=payload).json()["id"]
    _wait(client, f"/api/batch/{bid}", lambda s: s["state"] == "done", tries=300)
    runs = client.get("/api/runs").json()["runs"]
    assert runs[0]["id"] == bid and runs[0]["state"] == "done" and runs[0]["oncologists_saved"] > 0
    assert len(client.get(f"/api/runs/{bid}").json()["sites"]) == 2
    d = client.get("/api/directory").json()
    assert d["total"] > 0 and "Sunrise A" in d["facets"]["hospitals"]
    only_a = client.get("/api/directory", params={"hospital": "Sunrise A"}).json()
    assert 0 < only_a["total"] < d["total"] + 1 and all(x["source_hospital"] == "Sunrise A" for x in only_a["doctors"])
    assert client.get("/api/directory", params={"has_email": "true"}).json()["total"] <= d["total"]
    assert client.get("/api/directory", params={"q": "zzzz-nothing"}).json()["total"] == 0
    dash = client.get("/api/dashboard").json()
    assert dash["totals"]["oncologists"] == d["total"] and dash["hospitals"]
    assert client.get("/api/directory/download.xlsx").content[:2] == b"PK"
    assert "Sunrise" in client.get("/api/directory/download.csv").text
    assert client.get(f"/api/runs/{bid}/download.xlsx").content[:2] == b"PK"
    # app restart while a batch was going: it shows as interrupted and can be resumed
    web.store().update_run(bid, state="running")
    web.BATCHES.pop(bid)
    web._STORE = None
    assert client.get(f"/api/runs/{bid}").json()["state"] == "interrupted"
    r = client.post(f"/api/runs/{bid}/resume").json()
    assert r["id"] == bid
    st = _wait(client, f"/api/batch/{bid}", lambda s: s["state"] == "done", tries=300)
    assert {s["status"] for s in st["sites"]} == {"resumed"}  # finished hospitals were not crawled again
    # delete
    assert client.delete(f"/api/runs/{bid}").json()["ok"]
    assert all(x["id"] != bid for x in client.get("/api/runs").json()["runs"])
    web._STORE = None


def test_batch_crawls_several_sites_at_once(site_url, waf_url, tmp_path):
    """Two different websites run at the same time; rows on the same website stay one after another."""
    import threading
    from oncoscraper.batch import run_batch
    from oncoscraper.config import ScrapeConfig
    from oncoscraper.sources import Source
    srcs = [Source(row=2, hospital="Site A", urls=[f"{site_url}/specialities/cancer-care"]),
            Source(row=3, hospital="Site A2", urls=[f"{site_url}/faculty/radiation-oncology"]),
            Source(row=4, hospital="Site B", urls=[f"{waf_url}/doctor/dr-meenu-walia"])]
    live, seen_together = {}, []

    def progress(e):
        act = live.get("active") or {}
        if len(act) > 1:
            seen_together.append({s.hospital for _, s in act.values()})

    out = run_batch(srcs, tmp_path / "par.xlsx", ScrapeConfig(delay_seconds=0.2, render="never", use_sitemaps=False,
                                                              allow_private_hosts=True),
                    per_site_pages=4, resume=False, progress=progress, live=live, workers=4)
    assert out.is_file()
    assert seen_together and all(not {"Site A", "Site A2"} <= g for g in seen_together)
    assert any("Site B" in g for g in seen_together)
