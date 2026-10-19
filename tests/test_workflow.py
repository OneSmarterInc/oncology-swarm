from copy import deepcopy
from dataclasses import replace
from datetime import timedelta
from io import BytesIO
import threading
from concurrent.futures import ThreadPoolExecutor
import pytest
from openpyxl import load_workbook
from sqlalchemy import select, update, func
from app.cli import SEEDS
from app.db import Database
from app.models import Job, Proposal, Snapshot, Doctor, Control, utcnow
from app.service import Service, Conflict, BudgetExceeded
from app.worker import Coordinator
from app.export import export_xlsx, safe_cell
from app.validation import validate_proposal


def run_demo(service):
    ids = [service.enqueue(url)[0] for url in SEEDS]
    Coordinator(service).run(once=True)
    return ids


def test_end_to_end_manual_approval_export_and_restart(stack):
    settings, db, service = stack
    ids = run_demo(service)
    with db.session() as s:
        assert [s.get(Job, id).state for id in ids] == ['review', 'review']
        proposals = s.scalars(select(Proposal)).all()
        assert all(not p.validation_errors for p in proposals)
        assert s.get(Control, 1).requests_used == 0
        assert s.scalar(select(func.count()).select_from(Doctor)) == 0
    before = load_workbook(BytesIO(export_xlsx(db)))
    assert before['Doctors'].max_row == 1
    service.review(ids[0], 1, 'Test reviewer', 'Checked the synthetic source', True)
    db.engine.dispose()
    restarted = Database(settings.database_url)
    with restarted.session() as s:
        assert s.get(Job, ids[0]).state == 'approved'
        assert s.get(Job, ids[1]).state == 'review'
    wb = load_workbook(BytesIO(export_xlsx(restarted)))
    assert wb['Doctors'].max_row == 2
    assert wb['Doctors']['B2'].value == 'Dr Asha Rao'
    assert wb['Doctors']['E2'].value is True
    assert wb['Contacts']['D2'].value == 'institution_central'
    assert wb['Registration']['D2'].value == 'provider_claimed'
    assert wb['Evidence'].max_row > 5
    assert wb['Sources']['B2'].value == SEEDS[0]
    restarted.engine.dispose()


def test_enqueue_idempotency_and_durable_job_cap(stack):
    settings, db, service = stack
    first, created = service.enqueue(SEEDS[0] + '#top')
    assert created
    assert service.enqueue(SEEDS[0]) == (first, False)
    limited = Service(db, replace(settings, max_jobs=1))
    with pytest.raises(BudgetExceeded): limited.enqueue(SEEDS[1])
    with db.session() as s: assert s.get(Control, 1).jobs_created == 1


def test_concurrent_enqueue_and_claim_are_unique(stack):
    settings, db, service = stack
    with ThreadPoolExecutor(max_workers=4) as pool:
        result = list(pool.map(lambda _: service.enqueue(SEEDS[0]), range(12)))
    assert len({r[0] for r in result}) == 1
    assert sum(r[1] for r in result) == 1
    with ThreadPoolExecutor(max_workers=4) as pool:
        claims = list(pool.map(lambda i: service.claim('owner'+str(i)), range(4)))
    assert len([x for x in claims if x]) == 1


def test_pause_budget_and_request_count_atomic(stack):
    settings, db, service = stack
    service.enqueue(SEEDS[0])
    with db.session.begin() as s: s.get(Control, 1).paused = True
    assert service.claim('one') is None
    with db.session.begin() as s: s.get(Control, 1).paused = False
    limited = Service(db, replace(settings, max_requests=3))
    def reserve(_):
        try: limited.reserve_request('https://official.example/'); return True
        except BudgetExceeded: return False
    with ThreadPoolExecutor(max_workers=4) as pool: outcomes = list(pool.map(reserve, range(10)))
    assert sum(outcomes) == 3
    with db.session() as s: assert s.get(Control, 1).requests_used == 3


def test_lease_expiry_recovery_heartbeat_and_stale_fencing(stack):
    settings, db, service = stack
    id, _ = service.enqueue(SEEDS[0])
    assert service.claim('lost')[0] == id
    assert service.heartbeat(id, 'lost')
    with db.session.begin() as s: s.get(Job, id).lease_expires_at = utcnow() - timedelta(seconds=1)
    assert not service.heartbeat(id, 'lost')
    assert service.claim('new')[0] == id
    from app.fetcher import SafeFetcher
    from app.extractor import extract
    fetched = SafeFetcher(settings, service.reserve_request).fetch(SEEDS[0])
    assert not service.finish(id, 'lost', fetched, extract(fetched))
    service.fail(id, 'lost', 'stale error')
    assert service.finish(id, 'new', fetched, extract(fetched))
    with db.session() as s:
        assert s.get(Job, id).state == 'review'
        assert s.get(Job, id).attempts == 2
        assert s.scalar(select(func.count()).select_from(Proposal)) == 1


def test_retry_backoff_limit_and_manual_retry(stack):
    settings, db, service = stack
    id, _ = service.enqueue(SEEDS[0])
    for attempt in range(settings.max_attempts):
        assert service.claim('owner')
        service.fail(id, 'owner', 'temporary server error', retryable=True)
        with db.session.begin() as s:
            j = s.get(Job, id)
            assert j.state == ('failed' if attempt == settings.max_attempts - 1 else 'retry')
            j.available_at = utcnow() - timedelta(seconds=1)
    service.retry_failed(id)
    assert service.claim('new')


def test_validation_blocks_guessed_contacts_source_tampering_and_license_claims(stack):
    settings, db, service = stack
    id = run_demo(service)[0]
    with db.session() as s:
        p = s.scalar(select(Proposal).where(Proposal.job_id == id)); original = deepcopy(p.data)
        snap = s.scalar(select(Snapshot).where(Snapshot.job_id == id))
        forged = deepcopy(original); forged['contacts'][0]['value'] = 'guessed@example.com'
        assert validate_proposal(forged, snap)
        forged = deepcopy(original); forged['evidence'][0]['source_url'] = 'https://other.example/'
        assert validate_proposal(forged, snap)
        forged = deepcopy(original); forged['synthetic'] = False
        assert validate_proposal(forged, snap)
        forged = deepcopy(original); forged['registration'][0]['status'] = 'current_independently_verified'
        assert validate_proposal(forged, snap)
    assert service.edit(id, forged, 1)
    with pytest.raises(Conflict): service.review(id, 2, 'Human', 'I agree', True)
    assert not service.edit(id, original, 2)
    with pytest.raises(Conflict): service.review(id, 2, 'Human', 'stale', True)
    service.review(id, 3, 'Human', 'Supported', True)


def test_duplicate_detection_requires_human_acknowledgment(stack):
    settings, db, service = stack
    ids = run_demo(service)
    # Simulate a second immutable source genuinely supporting the same name.
    with db.session.begin() as s:
        p1 = s.scalar(select(Proposal).where(Proposal.job_id == ids[0]))
        p2 = s.scalar(select(Proposal).where(Proposal.job_id == ids[1]))
        source = s.scalar(select(Snapshot).where(Snapshot.job_id == ids[1]))
        first_source = s.scalar(select(Snapshot).where(Snapshot.job_id == ids[0]))
        source.text = first_source.text
        data = deepcopy(p1.data)
        for e in data['evidence']:
            e['source_url'], e['retrieved_at'] = source.url, source.retrieved_at
        p2.data = data
    with pytest.raises(Conflict, match='Possible duplicate'):
        service.review(ids[0], 1, 'Human', 'Checked', True)
    service.review(ids[0], 1, 'Human', 'Same name, separate source record', True, True)
    service.review(ids[1], 1, 'Human', 'Keeping separate intentionally', True, True)
    with db.session() as s: assert s.scalar(select(func.count()).select_from(Doctor)) == 2


def test_rejection_is_not_exported_or_silently_retried(stack):
    settings, db, service = stack
    id = run_demo(service)[0]
    service.review(id, 1, 'Human', 'Insufficient evidence', False)
    with pytest.raises(Conflict): service.retry_failed(id)
    assert load_workbook(BytesIO(export_xlsx(db)))['Doctors'].max_row == 1


def test_excel_formula_injection_is_escaped():
    for value in ['=HYPERLINK("https://bad")', '+1234', '-2+3', '@SUM(1)', '\t=1+1', '  =1+1']:
        assert safe_cell(value).startswith("'")
    assert safe_cell('Dr Test') == 'Dr Test'


def test_scope_evidence_must_label_the_exact_contact(stack):
    settings, db, service = stack
    from app.fetcher import FetchResult, visible_text
    from app.extractor import extract
    html = '<p>Doctor name: Dr Test Doctor</p><p>Hospital central email: central@hospital.example</p><p>No doctor email is published.</p>'
    fetched = FetchResult('https://hospital.example/profile', visible_text(html), html, '2026-10-19T00:00:00Z', 'a'*64, False)
    data = extract(fetched)
    assert not validate_proposal(data, fetched)
    data['contacts'][0]['scope'] = 'doctor'
    scope = next(e for e in data['evidence'] if e['field'] == 'contacts.0.scope')
    scope['value'], scope['support_text'] = 'doctor', 'No doctor email is published.'
    assert any('scope' in e for e in validate_proposal(data, fetched))


@pytest.mark.parametrize('first_action', ['edit', 'approve'])
def test_edit_and_approval_share_serialization_lock(stack, monkeypatch, first_action):
    settings, db, service = stack
    id = run_demo(service)[0]
    with db.session() as s: data = deepcopy(s.scalar(select(Proposal).where(Proposal.job_id == id)).data)
    import app.service as service_module
    original = service_module.validate_proposal
    entered, release = threading.Event(), threading.Event()
    def pause_first(*args):
        if threading.current_thread().name.startswith('first'):
            entered.set()
            assert release.wait(5)
        return original(*args)
    monkeypatch.setattr(service_module, 'validate_proposal', pause_first)
    results = {}
    def do(action, label):
        try:
            if action == 'edit': service.edit(id, data, 1)
            else: service.review(id, 1, 'Human', 'Verified', True)
            results[label] = 'ok'
        except Conflict: results[label] = 'conflict'
    a = threading.Thread(target=do, args=(first_action,'first'), name='first-operation')
    b = threading.Thread(target=do, args=('approve' if first_action == 'edit' else 'edit','second'), name='second-operation')
    a.start(); assert entered.wait(5); b.start(); release.set(); a.join(5); b.join(5)
    assert not a.is_alive() and not b.is_alive()
    assert results == {'first':'ok','second':'conflict'}
    with db.session() as s:
        p = s.scalar(select(Proposal).where(Proposal.job_id == id))
        assert p.version == 2
        assert s.get(Job,id).state == ('review' if first_action == 'edit' else 'approved')


def test_llm_budget_is_durable_and_separate(stack):
    settings, db, service = stack
    limited = Service(db, replace(settings, max_llm_calls=1))
    id1,_ = service.enqueue(SEEDS[0]); id2,_ = service.enqueue(SEEDS[1])
    service.claim('owner1'); service.claim('owner2')
    limited.reserve_llm_call(id1, 'owner1')
    with pytest.raises(BudgetExceeded): limited.reserve_llm_call(id2, 'owner2')
    with db.session() as s:
        c=s.get(Control,1)
        assert c.llm_calls_used == 1 and c.requests_used == 0


def test_ambiguous_llm_call_never_automatically_repeats_after_crash(stack):
    settings, db, service = stack
    id,_ = service.enqueue(SEEDS[0]); service.claim('lost')
    service.reserve_llm_call(id,'lost')
    with pytest.raises(Conflict): service.reserve_llm_call(id,'lost')
    with db.session.begin() as s: s.get(Job,id).lease_expires_at=utcnow()-timedelta(seconds=1)
    assert service.claim('recovery') is None
    with db.session() as s: assert s.get(Job,id).state=='failed'
    with pytest.raises(Conflict,match='billed'): service.retry_failed(id)
    service.retry_failed(id,True)
    assert service.claim('explicit-retry')[0]==id
    service.reserve_llm_call(id,'explicit-retry')
    with db.session() as s: assert s.get(Control,1).llm_calls_used==2
