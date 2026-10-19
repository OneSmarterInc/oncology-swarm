"""Independent adversarial regression cases for directory integration."""
from copy import deepcopy
from dataclasses import replace
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.cli import SEEDS
from app.export import export_xlsx
from app.extractor import extract
from app.fetcher import SafeFetcher
from app.main import create_app
from app.models import Audit, Control, Doctor, Job, Proposal, Snapshot
from app.service import Conflict
from app.worker import Coordinator


def _proposed(stack, duplicates=False):
    settings, db, service = stack
    ids = [service.enqueue(url)[0] for url in SEEDS]
    Coordinator(service).run(once=True)
    if duplicates:
        with db.session.begin() as s:
            a = s.scalar(select(Proposal).where(Proposal.job_id == ids[0]))
            b = s.scalar(select(Proposal).where(Proposal.job_id == ids[1]))
            first = s.scalar(select(Snapshot).where(Snapshot.job_id == ids[0]))
            second = s.scalar(select(Snapshot).where(Snapshot.job_id == ids[1]))
            data = deepcopy(a.data)
            second.text = first.text
            for evidence in data['evidence']:
                evidence['source_url'] = second.url
                evidence['retrieved_at'] = second.retrieved_at
            b.data = data
    return ids


def _approved(stack, duplicates=False):
    ids = _proposed(stack, duplicates)
    for job_id in ids:
        stack[2].review(job_id, 1, 'Audit reviewer', 'Checked source', True, duplicates)
    return ids


def _changed_refresh(stack, original, new_job):
    settings, db, service = stack
    assert service.claim('audit-refresh')[0] == new_job
    fetched = SafeFetcher(settings, service.reserve_request).fetch(SEEDS[0])
    fetched = replace(fetched, content_hash='f' * 64)
    assert service.finish(new_job, 'audit-refresh', fetched, extract(fetched))


def test_audit_invalid_doctor_shape_does_not_break_queue(stack):
    settings, db, service = stack
    job_id = service.enqueue(SEEDS[0])[0]
    service.claim('audit')
    fetched = SafeFetcher(settings, service.reserve_request).fetch(SEEDS[0])
    service.finish(job_id, 'audit', fetched, {'doctor': 'invalid model output'})
    with TestClient(create_app(settings, db), raise_server_exceptions=False,
                    headers={'Authorization': 'Bearer ' + settings.admin_token}) as client:
        response = client.get('/api/jobs')
        assert response.status_code == 200, response.text
        assert response.json()['items'][0]['validation_error_count'] > 0


def test_audit_edit_retains_actual_previous_data(stack):
    _, db, service = stack
    job_id = _proposed(stack)[0]
    with db.session() as s:
        previous = deepcopy(s.scalar(select(Proposal).where(Proposal.job_id == job_id)).data)
    changed = deepcopy(previous)
    changed['warnings'].append('Independent audit edit')
    service.edit(job_id, changed, 1)
    with db.session() as s:
        audit = s.scalar(select(Audit).where(Audit.job_id == job_id, Audit.action == 'edit'))
        assert audit.details['previous_data'] == previous
        assert audit.details['new_data'] == changed


def test_audit_changed_refresh_cannot_be_bypassed_by_old_correction(stack):
    _, db, service = stack
    original = _approved(stack)[0]
    fresh = service.refresh_job(original, 2, 'Audit reviewer', 'Refresh')['id']
    _changed_refresh(stack, original, fresh)
    try:
        service.correct(original, 2, 'Audit reviewer', 'Correct old snapshot')
        service.review(original, 3, 'Audit reviewer', 'Old snapshot checked', True, True)
    except Conflict:
        pass
    with db.session() as s:
        proposal = s.scalar(select(Proposal).where(Proposal.job_id == original))
        doctor = s.scalar(select(Doctor).where(Doctor.proposal_id == proposal.id))
        assert not doctor.active or doctor.needs_review, 'Old snapshot became exportable while changed refresh awaits review'


def test_audit_pending_refresh_cannot_bypass_merge_quarantine(stack):
    _, db, service = stack
    original, target = _approved(stack, duplicates=True)
    fresh = service.refresh_job(original, 2, 'Audit reviewer', 'Refresh')['id']
    try:
        service.merge(original, target, 2, 2, 'Audit reviewer', 'Same doctor')
    except Conflict:
        return
    _changed_refresh(stack, original, fresh)
    assert load_workbook(BytesIO(export_xlsx(db)))['Doctors'].max_row == 1, 'Merged target exports changed source without review'


def test_audit_model_usage_belongs_to_current_budget_run_after_retry(stack):
    _, db, service = stack
    job_id = service.enqueue(SEEDS[0])[0]
    service.claim('first')
    service.reserve_llm_call(job_id, 'first')
    service.fail(job_id, 'first', 'Provider failed')
    with db.session.begin() as s:
        s.get(Control, 1).paused = True
    assert service.reset_budgets('Audit reviewer', 'Next run')['run_number'] == 2
    service.retry_failed(job_id, True)
    with db.session.begin() as s:
        s.get(Control, 1).paused = False
    service.claim('second')
    service.reserve_llm_call(job_id, 'second')
    service.record_llm_usage(job_id, 'second', {
        'provider': 'audit', 'model': 'fake', 'input_tokens': 10,
        'output_tokens': 20, 'total_tokens': 30, 'estimated_cost_usd': 0.01,
    })
    metrics = service.metrics()
    assert metrics['llm_calls_reserved'] == 1
    assert metrics['llm_responses_recorded'] == 1
    assert metrics['estimated_cost_usd'] == 0.01


def test_audit_final_source_host_cannot_be_misattributed(stack):
    settings, db, original_service = stack
    from app.service import Service
    from app.fetcher import FetchResult
    from app.validation import validate_proposal
    live = Service(db, replace(settings, mode='live', allowed_domains=('first.example', 'second.example')))
    source = live.add_source('First Hospital', '', 'first.example', 'https://first.example/doctors',
                             'Audit reviewer', 'Only first host approved')
    job_id = live.enqueue('https://first.example/doctors/test')[0]
    live.claim('redirect')
    fetched = SafeFetcher(settings, original_service.reserve_request).fetch(SEEDS[0])
    # Worker transport currently permits redirects to any configured allowed domain.
    redirected = replace(fetched, url='https://second.example/doctors/test', synthetic=False)
    proposal = extract(redirected)
    assert not validate_proposal(proposal, redirected)
    try:
        done = live.finish(job_id, 'redirect', redirected, proposal)
    except (ValueError, Conflict):
        return
    if not done:
        return
    with db.session() as s:
        snapshot = s.scalar(select(Snapshot).where(Snapshot.job_id == job_id))
        assert snapshot.source_id != source['id'] or snapshot.source_type != 'official', 'Second host mislabeled as First Hospital approved official source'


def test_audit_merged_evidence_references_are_unambiguous(stack):
    _, db, service = stack
    original, target = _approved(stack, duplicates=True)
    service.merge(original, target, 2, 2, 'Audit reviewer', 'Same doctor')
    workbook = load_workbook(BytesIO(export_xlsx(db)))
    fields = list(workbook['Evidence'].values)
    headers, rows = fields[0], fields[1:]
    doctor_index, evidence_index = headers.index('doctor_id'), headers.index('evidence_id')
    seen = [(row[doctor_index], row[evidence_index]) for row in rows]
    assert len(seen) == len(set(seen)), 'Same exported doctor/evidence_id refers to multiple different source jobs; child sheets omit job_id'


def test_audit_concurrent_manual_retry_does_not_reset_new_paid_lease(stack, monkeypatch):
    import threading
    _, db, service = stack
    job_id = service.enqueue(SEEDS[0])[0]
    service.claim('original')
    service.reserve_llm_call(job_id, 'original')
    service.fail(job_id, 'original', 'Provider failed ambiguously')
    reached, release = threading.Event(), threading.Event()
    scalar = db.session.class_.scalar
    def paused_scalar(session, *args, **kwargs):
        if threading.current_thread().name == 'stale-retry' and not reached.is_set():
            reached.set()
            assert release.wait(5)
        return scalar(session, *args, **kwargs)
    monkeypatch.setattr(db.session.class_, 'scalar', paused_scalar)
    outcome = []
    def stale_retry():
        try:
            service.retry_failed(job_id, True)
            outcome.append('reset')
        except Conflict:
            outcome.append('conflict')
    thread = threading.Thread(target=stale_retry, name='stale-retry')
    thread.start()
    assert reached.wait(5)
    retry_done = threading.Event()
    second_outcome = []
    def second_retry():
        try:
            service.retry_failed(job_id, True)
            second_outcome.append('reset')
        except Conflict:
            second_outcome.append('conflict')
        finally:
            retry_done.set()
    second = threading.Thread(target=second_retry, name='second-retry')
    second.start()
    try:
        if retry_done.wait(0.3):
            # An unserialized stale reader must lose its conditional update.
            assert second_outcome == ['reset']
            assert service.claim('paid-worker')[0] == job_id
            service.reserve_llm_call(job_id, 'paid-worker')
            release.set()
        else:
            # An implementation may instead serialize retry before reading.
            release.set()
            thread.join(5)
            second.join(5)
            assert sorted(outcome + second_outcome) == ['conflict', 'reset']
            assert service.claim('paid-worker')[0] == job_id
            service.reserve_llm_call(job_id, 'paid-worker')
    finally:
        release.set()
        thread.join(5)
        second.join(5)
    assert not thread.is_alive() and not second.is_alive()
    with db.session() as s:
        job = s.get(Job, job_id)
        assert job.state == 'leased' and job.llm_call_reserved and job.lease_owner == 'paid-worker', f'Stale retry overwrote new paid lease; state={job.state}, reserved={job.llm_call_reserved}'


def test_audit_export_cannot_mix_approved_identity_with_unreviewed_evidence(stack, monkeypatch):
    import threading
    _, db, service = stack
    original = _approved(stack)[0]
    reached, release, edited = threading.Event(), threading.Event(), threading.Event()
    get = db.session.class_.get
    def paused_get(session, entity, *args, **kwargs):
        if threading.current_thread().name == 'audit-export' and entity is Proposal and not reached.is_set():
            reached.set()
            assert release.wait(5)
        return get(session, entity, *args, **kwargs)
    monkeypatch.setattr(db.session.class_, 'get', paused_get)
    exported, failures = [], []
    def export_thread():
        try:
            exported.append(export_xlsx(db))
        except Exception as exc:
            failures.append(exc)
    def edit_thread():
        try:
            service.correct(original, 2, 'Audit reviewer', 'Recheck evidence')
            with db.session() as s:
                data = deepcopy(s.scalar(select(Proposal).where(Proposal.job_id == original)).data)
            data['evidence'] = []
            service.edit(original, data, 3)
        except Exception as exc:
            failures.append(exc)
        finally:
            edited.set()
    exporting = threading.Thread(target=export_thread, name='audit-export')
    exporting.start()
    assert reached.wait(5)
    editing = threading.Thread(target=edit_thread, name='audit-edit')
    editing.start()
    # A shared lock may serialize the edit; otherwise export must use a snapshot.
    edited.wait(0.3)
    release.set()
    exporting.join(5)
    editing.join(5)
    assert not exporting.is_alive() and not editing.is_alive()
    assert not failures, failures
    workbook = load_workbook(BytesIO(exported[0]))
    doctors = list(workbook['Doctors'].values)
    evidence = list(workbook['Evidence'].values)
    name_column = doctors[0].index('name')
    for row in doctors[1:]:
        if row[name_column] == 'Dr Asha Rao':
            assert any(e[0] == row[0] and e[3] == 'doctor.name' for e in evidence[1:]), 'Export included old approved doctor but new unapproved empty evidence'


def test_audit_unmerge_rejects_inner_contributor_before_outer_split(stack):
    _, db, service = stack
    first, middle = _approved(stack, duplicates=True)
    with db.session.begin() as s:
        third = Job(url=SEEDS[0], host='demo.example.invalid', generation=1, run_number=1)
        s.add(third)
        s.flush()
        outer = third.id
    Coordinator(service).run(once=True)
    service.review(outer, 1, 'Audit reviewer', 'Same person third source', True, True)
    service.merge(first, middle, 2, 2, 'Audit reviewer', 'First into middle')
    service.merge(middle, outer, 3, 2, 'Audit reviewer', 'Middle group into outer')
    with pytest.raises(Conflict, match='outer|nested|group|contributor'):
        service.unmerge(first, outer, 3, 3, 'Audit reviewer', 'Wrong split order')
    # The outer group can split safely, then the inner group can split.
    service.unmerge(middle, outer, 4, 3, 'Audit reviewer', 'Split outer first')
    service.unmerge(first, middle, 3, 5, 'Audit reviewer', 'Then split inner')
    for job_id in (first, middle, outer):
        with db.session() as s:
            proposal = s.scalar(select(Proposal).where(Proposal.job_id == job_id))
            version = proposal.version
        service.correct(job_id, version, 'Audit reviewer', 'Re-review individual')


def test_audit_expired_discovery_cannot_enqueue_after_budget_reset(stack, monkeypatch):
    import threading
    from datetime import timedelta
    from app.discovery import DiscoveredProfile, DiscoveryResult
    from app.models import DiscoveryRun, utcnow
    from app.service import Service
    settings, db, _ = stack
    service = Service(db, replace(settings, mode='live', allowed_domains=('hospital.example',)))
    source = service.add_source('Test Hospital', '', 'hospital.example', 'https://hospital.example/doctors',
                                'Audit reviewer', 'Approved test source')
    reached, release = threading.Event(), threading.Event()
    def delayed_discover(*args, **kwargs):
        reached.set()
        assert release.wait(5)
        result = DiscoveryResult(seed_url=source['list_url'])
        result.profiles.append(DiscoveredProfile('https://hospital.example/doctors/a', source['list_url'], 'Dr A', 'Oncology'))
        return result
    monkeypatch.setattr('app.discovery.discover', delayed_discover)
    outcomes = []
    def discovery():
        try:
            service.discover_source(source['id'])
            outcomes.append('completed')
        except Conflict:
            outcomes.append('fenced')
    thread = threading.Thread(target=discovery, name='old-discovery')
    thread.start()
    assert reached.wait(5)
    with db.session.begin() as s:
        run = s.scalar(select(DiscoveryRun))
        run.lease_expires_at = utcnow() - timedelta(seconds=1)
        s.get(Control, 1).paused = True
    assert service.reset_budgets('Audit reviewer', 'Old discovery lease expired')['run_number'] == 2
    release.set()
    thread.join(5)
    assert not thread.is_alive()
    with db.session() as s:
        assert not s.scalars(select(Job)).all(), 'Expired old discovery enqueued jobs into a newly authorized budget run'
        assert s.scalar(select(DiscoveryRun)).state == 'failed'


def test_audit_reverted_source_can_recover_after_changed_refresh_rejected(stack):
    _, db, service = stack
    original = _approved(stack)[0]
    changed = service.refresh_job(original, 2, 'Audit reviewer', 'Check source')['id']
    _changed_refresh(stack, original, changed)
    service.review(changed, 1, 'Audit reviewer', 'Reject changed evidence', False)
    reverted = service.refresh_job(original, 2, 'Audit reviewer', 'Source reverted; check again')['id']
    Coordinator(service).run(once=True)
    with db.session() as s:
        job = s.get(Job, reverted)
        proposal = s.scalar(select(Proposal).where(Proposal.job_id == reverted))
    assert job.state == 'review' and proposal, 'Previously quarantined source was marked unchanged with no recovery/review path'
    service.review(reverted, proposal.version, 'Audit reviewer', 'Reverted current source checked', True, True)
    assert load_workbook(BytesIO(export_xlsx(db)))['Doctors'].max_row == 3


def test_audit_merged_child_evidence_references_resolve_exactly_once(stack):
    _, db, service = stack
    first, target = _approved(stack, duplicates=True)
    service.merge(first, target, 2, 2, 'Audit reviewer', 'Same doctor')
    workbook = load_workbook(BytesIO(export_xlsx(db)))
    evidence_rows = list(workbook['Evidence'].values)
    evidence_headers = evidence_rows[0]
    ids = [(row[evidence_headers.index('doctor_id')], row[evidence_headers.index('evidence_id')]) for row in evidence_rows[1:]]
    for sheet in ('Affiliations', 'Contacts', 'Registration'):
        rows = list(workbook[sheet].values)
        headers = rows[0]
        for row in rows[1:]:
            for ref in row[headers.index('evidence_ids')].split(', '):
                assert ids.count((row[headers.index('doctor_id')], ref)) == 1


def test_audit_orphan_evidence_cannot_export_unsupported_claim(stack):
    _, db, service = stack
    job_id = _proposed(stack)[0]
    with db.session() as s:
        proposal = s.scalar(select(Proposal).where(Proposal.job_id == job_id))
        data = deepcopy(proposal.data)
    forged = deepcopy(data['evidence'][0])
    forged.update(id='orphan-qualification', field='doctor.qualifications', value='Fellowship never present in snapshot')
    data['doctor']['qualifications'] = None
    data['evidence'].append(forged)
    errors = service.edit(job_id, data, 1)
    assert errors, 'Orphan evidence accepted an unsupported value just because field is null/missing from facts'
    with pytest.raises(Conflict, match='Validation'):
        service.review(job_id, 2, 'Audit reviewer', 'Must not approve fabricated evidence', True)
