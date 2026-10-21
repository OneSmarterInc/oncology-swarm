"""Hospital source lists: the standard template, reading (.xlsx/.csv) and validation.

Each row is one hospital (or one unit of a chain). The crawler stays on the
website of the URLs given in that row, so there is no separate allow-list to
maintain: the uploaded file is the allow-list.
"""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

# (key, header shown in Excel, required?, description, example)
TEMPLATE_COLUMNS = [
    ("hospital_name", "Hospital Name", "Required",
     "Hospital, cancer centre or the specific unit of a chain.", "Rajiv Gandhi Cancer Institute"),
    ("official_website", "Official Website", "Required if no page URL below",
     "Home page or domain. Used as the start point when no page URL is given.", "https://www.rgcirc.org"),
    ("oncology_page_urls", "Oncology Page URLs", "Recommended",
     "Page(s) listing the oncology / cancer-care doctors. Several URLs allowed: separate with a space or new line.",
     "https://www.rgcirc.org/find-a-doctor/"),
    ("doctors_list_urls", "Doctors List URLs", "Optional",
     "The general 'Find a Doctor' / 'Our Doctors' page(s). Used to find oncologists the oncology page misses.",
     "https://www.rgcirc.org/find-a-doctor/"),
    ("sample_profile_url", "Sample Doctor Profile URL", "Recommended",
     "Any ONE doctor's profile page on the site. The scraper learns the profile URL pattern from it.",
     "https://www.rgcirc.org/doctor/dr-vineet-talwar/"),
    ("profile_url_pattern", "Profile URL Pattern", "Optional",
     "Path that all profile URLs share, e.g. /doctor/. Leave blank: it is worked out from the sample profile.", ""),
    ("city", "City", "Optional", "City of the hospital or unit.", "New Delhi"),
    ("state", "State", "Optional", "State.", "Delhi"),
    ("hospital_type", "Hospital Type", "Optional",
     "private_chain, cancer_centre, government, trust or other.", "cancer_centre"),
    ("max_pages", "Max Pages", "Optional", "Page budget for this hospital. Blank = the run's default (150).", "150"),
    ("include_subdomains", "Include Subdomains", "Optional",
     "yes = also crawl other subdomains of the site (e.g. doctors.hospital.in). Default no.", "no"),
    ("use_browser", "Use Browser", "Optional",
     "auto (default) = real browser only when needed; always = every page in the browser; never = plain HTTP only.",
     "auto"),
    ("enabled", "Enabled", "Optional", "no = skip this row without deleting it. Default yes.", "yes"),
    ("notes", "Notes", "Optional", "Anything useful. Not used by the scraper.", ""),
]
KEYS = [c[0] for c in TEMPLATE_COLUMNS]

ALIASES = {
    "hospital_name": ["hospital_name", "hospital", "name", "institution", "hospital_or_unit", "centre", "center"],
    "official_website": ["official_website", "official_domain", "website", "domain", "homepage", "home_page", "site"],
    "oncology_page_urls": ["oncology_page_urls", "oncology_page_url", "oncology_urls", "oncology_url", "cancer_page_url",
                           "cancer_care_url", "start_url", "url"],
    "doctors_list_urls": ["doctors_list_urls", "doctors_list_url", "doctor_list_url", "find_a_doctor_url",
                          "doctors_page_url", "our_doctors_url"],
    "sample_profile_url": ["sample_profile_url", "sample_doctor_profile_url", "profile_url", "example_profile_url"],
    "profile_url_pattern": ["profile_url_pattern", "profile_pattern", "pattern"],
    "city": ["city", "city_or_network", "location"],
    "state": ["state"],
    "hospital_type": ["hospital_type", "type"],
    "max_pages": ["max_pages", "pages", "page_budget"],
    "include_subdomains": ["include_subdomains", "subdomains"],
    "use_browser": ["use_browser", "browser", "render"],
    "enabled": ["enabled", "active", "include", "run"],
    "notes": ["notes", "note", "comments", "remarks"],
}
_LOOKUP = {alias: key for key, names in ALIASES.items() for alias in names}
YES = {"yes", "y", "true", "1"}
NO = {"no", "n", "false", "0"}


def _norm_header(h) -> str:
    h = re.sub(r"\(.*?\)|\*", "", str(h or "")).strip().lower()
    return re.sub(r"[^a-z0-9]+", "_", h).strip("_")


@dataclass
class Source:
    row: int
    hospital: str
    urls: list[str]
    patterns: list[str] = field(default_factory=list)
    max_pages: int | None = None
    include_subdomains: bool = False
    render: str | None = None
    city: str = ""
    hospital_type: str = ""
    warnings: list[str] = field(default_factory=list)

    def as_dict(self) -> dict:
        return {"row": self.row, "hospital": self.hospital, "urls": self.urls, "patterns": self.patterns,
                "max_pages": self.max_pages, "include_subdomains": self.include_subdomains, "render": self.render,
                "city": self.city, "hospital_type": self.hospital_type, "warnings": self.warnings}


@dataclass
class Problem:
    row: int
    hospital: str
    problem: str


def _urls(cell: str) -> tuple[list[str], list[str]]:
    good, bad = [], []
    for token in re.split(r"[\s,;|]+", cell or ""):
        token = token.strip()
        if not token:
            continue
        if not re.match(r"https?://", token, re.I):
            if re.match(r"^(?:www\.)?[a-z0-9-]+(?:\.[a-z0-9-]+)+(?:/.*)?$", token, re.I):
                token = "https://" + token
            else:
                bad.append(token)
                continue
        p = urlsplit(token)
        if not p.hostname or "." not in p.hostname:
            bad.append(token)
        elif p.path.lower().endswith(".pdf"):
            bad.append(token + " (PDF not supported yet)")
        elif token not in good:
            good.append(token)
    return good, bad


def pattern_from_profile(url: str) -> str | None:
    """'/doctor/dr-meenu-walia' -> '/doctor/';  '/dr-senthil-rajappa/' -> '/dr-'."""
    path = urlsplit(url).path.rstrip("/")
    if not path:
        return None
    parent, _, slug = path.rpartition("/")
    if parent:
        return parent + "/"
    m = re.match(r"^((?:dr|doctor|prof)[-_.])", slug, re.I)
    return "/" + m.group(1) if m else None


def _rows(path: Path) -> list[dict]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["Hospitals"] if "Hospitals" in wb.sheetnames else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = next(it, None) or []
        rows = [dict(zip(header, r)) for r in it]
        wb.close()
    elif path.suffix.lower() in (".csv", ".txt"):
        with open(path, newline="", encoding="utf-8-sig") as fh:
            rows = list(csv.DictReader(fh))
    else:
        raise ValueError("Hospital list must be .xlsx or .csv")
    out = []
    for raw in rows:
        row: dict[str, str] = {}
        for k, v in raw.items():
            key = _LOOKUP.get(_norm_header(k))
            if key and v not in (None, ""):
                text = str(v).strip()
                row[key] = (row[key] + " " + text) if key in row and key.endswith("urls") else row.get(key) or text
        out.append(row)
    return out


def read_sources(path: str | Path) -> tuple[list[Source], list[Problem]]:
    """Read and check a hospital list. Returns usable sources and a list of row problems."""
    sources: list[Source] = []
    problems: list[Problem] = []
    for i, row in enumerate(_rows(Path(path)), start=2):  # row 1 is the header in Excel
        if not any(row.values()):
            continue
        name = row.get("hospital_name", "")
        if (row.get("enabled", "").lower() in NO):
            problems.append(Problem(i, name, "skipped: Enabled = no"))
            continue
        onco, bad1 = _urls(row.get("oncology_page_urls", ""))
        lists, bad2 = _urls(row.get("doctors_list_urls", ""))
        sample, bad3 = _urls(row.get("sample_profile_url", ""))
        home, bad4 = _urls(row.get("official_website", ""))
        warnings = [f"ignored invalid URL: {b}" for b in bad1 + bad2 + bad3 + bad4]
        urls = list(dict.fromkeys(onco + lists + sample[:1]))
        if not urls and home:
            urls = [home[0].rstrip("/") + "/"]
            warnings.append("no page URL given; starting from the home page (uses more of the page budget)")
        if not urls:
            problems.append(Problem(i, name, "no usable URL: fill Oncology Page URLs, Doctors List URLs or Official Website"))
            continue
        if not name:
            name = urlsplit(urls[0]).hostname or f"row {i}"
            warnings.append("Hospital Name missing; using the website name")
        patterns: list[str] = []
        pat = row.get("profile_url_pattern", "")
        if pat:
            patterns.append(pat if pat.startswith("/") else "/" + pat)
        elif sample:
            derived = pattern_from_profile(sample[0])
            if derived:
                patterns.append(derived)
        hosts = {(urlsplit(u).hostname or "").removeprefix("www.") for u in urls}
        if len(hosts) > 1:
            warnings.append("URLs span several websites: " + ", ".join(sorted(h for h in hosts if h)))
        mp = str(row.get("max_pages", "")).strip()
        try:
            max_pages = int(float(mp)) if mp else None
        except ValueError:
            max_pages = None
            warnings.append(f"Max Pages '{mp}' is not a number; using the default")
        browser = row.get("use_browser", "").lower()
        render = browser if browser in ("auto", "always", "never") else None
        if browser and not render:
            warnings.append(f"Use Browser '{browser}' not understood; using auto")
        sources.append(Source(i, name, urls, patterns, max_pages,
                              row.get("include_subdomains", "").lower() in YES, render,
                              row.get("city", ""), row.get("hospital_type", ""), warnings))
    return sources, problems


def write_template(path: str | Path, rows: list[dict] | None = None) -> Path:
    """Write the standard hospital-list workbook (Hospitals + Instructions sheets)."""
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Hospitals"
    ws.append([c[1] for c in TEMPLATE_COLUMNS])
    head_fill = PatternFill("solid", fgColor="164E63")
    for cell, col in zip(ws[1], TEMPLATE_COLUMNS):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill if col[2].startswith(("Required", "Recommended")) else PatternFill("solid", fgColor="4B6B73")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.comment = Comment(f"{col[2]}. {col[3]}", "OncoScraper")
    examples = rows if rows is not None else [{c[0]: c[4] for c in TEMPLATE_COLUMNS}]
    for r in examples:
        ws.append([r.get(k, "") for k in KEYS])
    widths = {"hospital_name": 34, "official_website": 30, "oncology_page_urls": 55, "doctors_list_urls": 45,
              "sample_profile_url": 55, "profile_url_pattern": 18, "notes": 50}
    for idx, key in enumerate(KEYS, start=1):
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = widths.get(key, 14)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"
    last = max(ws.max_row, 500)
    for key, options in (("hospital_type", "private_chain,cancer_centre,government,trust,other"),
                         ("include_subdomains", "yes,no"), ("use_browser", "auto,always,never"), ("enabled", "yes,no")):
        col = ws.cell(1, KEYS.index(key) + 1).column_letter
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
        dv.add(f"{col}2:{col}{last}")
        ws.add_data_validation(dv)
    ins = wb.create_sheet("Instructions")
    ins.append(["Column", "Required?", "What to put", "Example"])
    for cell in ins[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
    for _key, header, req, desc, ex in TEMPLATE_COLUMNS:
        ins.append([header, req, desc, ex])
    ins.append([])
    for line in [
        "One row per hospital, or per unit of a chain if each unit has its own doctor page.",
        "Fill at least one of: Oncology Page URLs, Doctors List URLs, Official Website.",
        "Best results: the oncology/cancer-care doctors page plus one sample doctor profile URL.",
        "Use the hospital's own website only. Aggregators (Practo, Lybrate, JustDial) are not official sources.",
        "The crawler only visits the websites in each row; there is no separate allow-list to edit.",
        "robots.txt is always obeyed; sites that ask crawlers to stay out are skipped and reported.",
        "PDF doctor lists are not read yet; put an HTML page instead where one exists.",
        "Upload this file in the web UI (Batch section) or run: python -m oncoscraper batch <file>.xlsx",
    ]:
        ins.append(["", "", line])
    for col, w in zip("ABCD", (26, 28, 95, 48)):
        ins.column_dimensions[col].width = w
    for row in ins.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
