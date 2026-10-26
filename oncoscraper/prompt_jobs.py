"""Durable prompt -> hospital discovery -> existing crawler -> filtered Excel workflow."""
from __future__ import annotations

import threading
import time
import uuid
from dataclasses import asdict
from pathlib import Path
from typing import Literal

from pydantic import BaseModel, ConfigDict, Field, field_validator

from .batch import run_batch
from .config import ScrapeConfig
from .crawl import Crawler, CrawlResult, PageLog
from .discovery_services import (
    CachedFetcher, CachedSearch, ConfigurationError, ModelRouter, Paused, ServiceWait, settings,
)
from .export import to_xlsx
from .fetch import Fetcher
from .models import DoctorRecord
from .prompt_discovery import Limits, Report
from .site_discovery import Hospital, discover_hospitals, read_request, sources_from_hospitals
from .verified_discovery import discover_verified, filter_records, sources_from_report

ACTIVE = {"queued", "running", "waiting", "pausing", "stopping"}
TERMINAL = {"done", "partial", "clarify", "unsupported", "no_candidates", "sources_ready"}
CONTACT = ("phone", "email", "contact_phone", "contact_email")


class PromptIn(BaseModel):
    model_config = ConfigDict(extra="forbid")
    prompt: str = Field(min_length=1, max_length=6000)
    max_hospitals: int = Field(10, ge=1, le=100)
    discovery_pages: int = Field(60, ge=1, le=100)
    searches: int = Field(12, ge=1, le=30)
    pages_per_hospital: int = Field(150, ge=1, le=1000)
    render: Literal["never", "auto", "always"] = "auto"
    auto_scrape: bool = True
    # "sites" proves each hospital by crawling it; "evidence" asks a model to judge pages.
    engine: Literal["sites", "evidence"] = "sites"
    site_pages: int = Field(20, ge=2, le=80)
    site_seconds: int = Field(60, ge=10, le=300)  # wall clock one site may take while being checked
    site_workers: int = Field(6, ge=1, le=12)     # sites checked at the same time

    @field_validator("prompt")
    @classmethod
    def nonempty(cls, value):
        if not value.strip():
            raise ValueError("Describe the hospitals or oncologists you need")
        return value.strip()


class PromptManager:
    def __init__(self, db, history, out_dir):
        self.db, self.history, self.out_dir = db, history, Path(out_dir)
        self.lock = threading.RLock()
        self.wake = threading.Event()
        self.closed = threading.Event()
        self.thread = None
        self.current = None
        self.stop = threading.Event()
        # A new process cannot still own an in-flight HTTP/browser operation.
        for job in self.db.jobs():
            if job["state"] in {"running", "pausing", "stopping"}:
                job["state"] = "interrupted"
                self.db.save(job)
                self.history.update_run(job["id"], state="interrupted")

    def start_worker(self):
        with self.lock:
            if self.thread is None or not self.thread.is_alive():
                self.thread = threading.Thread(target=self._loop, daemon=True, name="prompt-discovery")
                self.thread.start()
        self.wake.set()

    def close(self):
        self.closed.set()
        self.stop.set()
        self.wake.set()

    def create(self, body, start=True):
        config = settings()
        if body.engine == "evidence" and not (config["ollama_configured"] or config["groq_configured"]
                                              or config["gemini_configured"]):
            raise ConfigurationError("'AI judges each page' needs Ollama or an explicitly selected cloud provider "
                                     "in .env. 'Crawl each site' needs only SearXNG")
        job = {"id": uuid.uuid4().hex[:10], "state": "queued", "phase": "queued", "started": time.time(),
               "finished": None, "body": body.model_dump(), "report": {}, "rows": [], "progress": {},
               "error": None, "retry_at": None, "xlsx": None, "stats": {}}
        self.history.create_run(job["id"], "prompt", body.prompt[:120], body.model_dump())
        self.history.update_run(job["id"], state="queued")
        self.db.save(job)
        if start:
            self.start_worker()
        return job

    def update(self, job_id, **values):
        with self.lock:
            job = self.db.get(job_id)
            if job is None:
                return
            job.update(values)
            self.db.save(job)
            hist = {k: values[k] for k in ("state", "finished", "error", "xlsx", "stats") if k in values}
            if hist:
                self.history.update_run(job_id, **hist)

    def action(self, job_id, action):
        with self.lock:
            job = self.db.get(job_id)
            if not job:
                raise KeyError("Prompt run not found")
            state = job["state"]
            if action in {"pause", "stop"}:
                if state not in ACTIVE:
                    raise ValueError(f"Cannot {action} a run that is {state}")
                active = self.current == job_id
                next_state = (("pausing" if action == "pause" else "stopping") if active
                              else ("paused" if action == "pause" else "stopped"))
                self.update(job_id, state=next_state, retry_at=None)
                if active:
                    self.stop.set()
            elif action == "resume":
                if state not in {"paused", "stopped", "interrupted", "waiting", "failed"}:
                    raise ValueError(f"Cannot resume a run that is {state}")
                self.update(job_id, state="queued", error=None, retry_at=None, finished=None)
                self.start_worker()
            else:
                raise ValueError("Unknown action")
        self.wake.set()
        return self.db.get(job_id)

    def _loop(self):
        while not self.closed.is_set():
            with self.lock:
                pending = [j for j in reversed(self.db.jobs()) if j["state"] == "queued" or
                           (j["state"] == "waiting" and (j.get("retry_at") or 0) <= time.time())]
                job = pending[0] if pending else None
                if job:
                    self.current = job["id"]
                    self.stop = threading.Event()
                    self.update(job["id"], state="running", error=None, retry_at=None)
            if job:
                self.execute(job["id"])
                with self.lock:
                    self.current = None
            else:
                self.wake.wait(2)
                self.wake.clear()

    def _discover_sites(self, body, search, progress, control):
        """Find hospitals and their doctor pages by crawling each site."""
        request = read_request(body.prompt)
        want = min(request.count, body.max_hospitals)
        report = Report(prompt=body.prompt)
        criteria = [{"id": "institution_type", "requirement": f"{request.kind} hospital",
                     "prompt_quote": request.kind, "applies_to": "institution"}] if request.kind else []
        if request.location:
            criteria.append({"id": "location", "requirement": f"Located in {request.location}",
                             "prompt_quote": request.location, "applies_to": "institution"})
        report.plan = {"objective": body.prompt, "target": request.target, "criteria": criteria,
                       "requested_institutions": request.count, "explanation": request.summary(),
                       "request": request.as_dict(), "hospitals": []}
        if request.count > body.max_hospitals:
            report.issues.append(f"Requested {request.count} hospitals but this run is limited to "
                                 f"{body.max_hospitals}. Raise 'Maximum hospitals' to {request.count}")
        if request.specialties:
            report.issues.append("Keeping only these specialities in the contact export: "
                                 + ", ".join(request.specialties))
        if request.unfiltered_conditions:
            report.issues.append("This run verifies hospitals and doctor pages. Conditions about individual "
                                 "doctors (email required, experience, language, gender) were not applied; "
                                 "switch the engine to 'AI judges each page' for those")

        def make_crawler():
            # Checking a site is a quick look, not the real extraction: no browser,
            # short timeouts, no retry. The accepted pages are crawled properly later.
            return Crawler(ScrapeConfig(render="never", strict_public_network=True, max_bytes=2_000_000,
                                        retries=0, max_pages=body.site_pages, stale_pages=0,
                                        timeout_seconds=10.0, delay_seconds=0.4, use_sitemaps=False,
                                        oncology_only=request.target == "oncologists"),
                           stop_event=self.stop)

        def site_progress(event):
            """Keep the counters on screen moving while sites are being read."""
            report.searches = max(report.searches, event.get("searches", 0))
            report.pages = max(report.pages, event.get("pages", 0))
            checked, sites = event.get("checked"), event.get("sites")
            if checked is not None and sites:
                event["current"] = f"{event['current']} ({checked} of {sites} sites checked)"
            event["report"] = asdict(report)
            progress(event)

        hospitals, rejected, notes = discover_hospitals(
            search, make_crawler, want=want, kind=request.kind, location=request.location,
            target=request.target, names=request.names, max_searches=body.searches,
            workers=body.site_workers, site_seconds=body.site_seconds,
            control=control, progress=site_progress)
        report.plan["hospitals"] = [h.as_dict() for h in hospitals]
        report.pages = max(report.pages, sum(h.pages_read for h in hospitals))
        for hospital in hospitals:
            found = hospital.oncologists_found if request.target == "oncologists" else hospital.doctors_found
            for url in hospital.doctor_urls:
                report.candidates.append({
                    "url": url, "institution": hospital.name, "accepted": True, "extracted_staff": found,
                    "reason": f"{found} {request.target} read from this site by the extractor",
                    "evidence_sources": [{"criterion_id": k, "quote": v, "source_url": hospital.homepage}
                                         for k, v in hospital.evidence.items()]})
        report.attempts = [{"url": f"https://{r.host}/", "institution": r.name, "accepted": False,
                            "reason": r.reason, "evidence_sources": []} for r in rejected]
        report.issues.extend(notes)
        report.issues = list(dict.fromkeys(report.issues))
        report.status = "candidates_found" if report.candidates else "no_candidates"
        return report

    def execute(self, job_id):
        """Public for deterministic integration tests; production uses one worker."""
        job = self.db.get(job_id)
        body = PromptIn(**job["body"])
        model = search = fetcher = None

        def control():
            if self.closed.is_set() or self.stop.is_set():
                raise Paused()

        def progress(event):
            phase = event.pop("phase", "checking_pages")
            changes = {"phase": phase, "progress": event}
            if "report" in event:
                changes["report"] = event.pop("report")
            self.update(job_id, **changes)
            control()  # Persist the operation that just finished even if Pause was clicked.

        def fail(message):
            latest = self.db.get(job_id)
            report = latest.get("report") or asdict(Report(prompt=body.prompt))
            planning = latest.get("phase") in {"planning", "planning_failed"}
            report["status"] = "planning_failed" if planning else "failed"
            if message not in report.setdefault("issues", []):
                report["issues"].append(message)
            self.update(job_id, state="failed", phase="planning_failed" if planning else latest["phase"],
                        report=report, error=message, finished=time.time())

        try:
            self.update(job_id, state="running")
            saved = self.db.cached(job_id, "discovery_report")
            if saved:
                report = Report(**saved)
            else:
                config = settings()
                search = CachedSearch(self.db, job_id, config["search_url"], control)
                if body.engine == "sites":
                    report = self._discover_sites(body, search, progress, control)
                else:
                    model = ModelRouter(self.db, job_id, control)
                    fetcher = CachedFetcher(Fetcher(ScrapeConfig(render=body.render, strict_public_network=True,
                                                                 max_bytes=2_000_000, retries=1)),
                                            self.db, job_id, control)
                    report = discover_verified(
                        body.prompt, model, search, fetcher,
                        Limits(searches=body.searches, pages=body.discovery_pages, institutions=body.max_hospitals,
                               rounds=2), control, progress,
                        load_state=lambda: self.db.cached(job_id, "discovery_state_v1"),
                        save_state=lambda state: self.db.cache(job_id, "discovery_state_v1", state))
                self.db.cache(job_id, "discovery_report", asdict(report))
            self.update(job_id, report=asdict(report))
            if report.status != "candidates_found":
                self.update(job_id, state=report.status, phase=report.status, finished=time.time())
                return
            sources = (sources_from_hospitals([Hospital(**h) for h in report.plan["hospitals"]])
                       if body.engine == "sites" else sources_from_report(report))
            if not body.auto_scrape:
                self.update(job_id, state="sources_ready", phase="sources_ready", finished=time.time())
                return
            target = report.plan.get("target", "oncologists")
            request = report.plan.get("request") or {}
            extracted = self.db.cached(job_id, "extracted")
            if extracted is None:
                rows = {s.row: {"sequence": i, "row": s.row, "hospital": s.hospital, "status": "queued",
                                "pages": 0, "oncologists": 0} for i, s in enumerate(sources, 1)}
                self.update(job_id, phase="scraping", rows=list(rows.values()))
                live, events_lock = {}, threading.RLock()

                def crawl_event(e):
                    with events_lock:
                        row = rows.get(e.get("row"))
                        if row is not None:
                            if e["event"] == "site_start":
                                row["status"] = "scraping"
                            elif e["event"] == "site_done":
                                row.update({k: v for k, v in e.items() if k != "event"})
                        for crawler, src in list((live.get("active") or {}).values()):
                            if src.row in rows:
                                rows[src.row]["pages"] = len(crawler.result_pages)
                        done = sum(r["status"] not in {"queued", "scraping"} for r in rows.values())
                        self.update(job_id, rows=list(rows.values()), progress={"completed": done, "total": len(rows)})
                    # No control() here: this runs inside crawl workers. run_batch watches self.stop itself.

                self.out_dir.mkdir(parents=True, exist_ok=True)
                run_batch(sources, self.out_dir / f"prompt-{job_id}-raw.xlsx",
                          ScrapeConfig(render=body.render, strict_public_network=True,
                                       oncology_only=target == "oncologists"),
                          body.pages_per_hospital, resume=True, progress=crawl_event, stop_event=self.stop,
                          live=live, workers=2)
                control()
                if len(live["summary"]) != len(sources):
                    raise RuntimeError("Some hospital workers did not finish. Resume to retry")
                extracted = {"doctors": [d.to_full_dict() for d in live["merger"].records],
                             "pages": [asdict(p) for p in live["pages"]], "sites": live["summary"]}
                self.db.cache(job_id, "extracted", extracted)
            control()
            docs = [DoctorRecord.from_full_dict(d) for d in extracted["doctors"]]
            if target == "oncologists":
                docs = [d for d in docs if d.is_oncology]
            if request.get("specialties"):
                wanted = set(request["specialties"])
                before = len(docs)
                docs = [d for d in docs if wanted & set(d.get("sub_specialties") or [])]
                report.issues.append(f"Speciality filter kept {len(docs)} of {before} doctors")
            location = request.get("location") or ""
            if location:
                # Hospital chains list doctors from many cities; drop the ones the site places elsewhere.
                from .vocab import CITY_CANON
                want = CITY_CANON.get(location, location).casefold()
                keep = [d for d in docs if not d.get("city")
                        or CITY_CANON.get(d.get("city"), d.get("city")).casefold() == want]
                if len(keep) < len(docs):
                    report.issues.append(f"{len(docs) - len(keep)} doctors the hospital site places outside "
                                         f"{location} were left out")
                docs = keep
            self.update(job_id, phase="filtering")
            filtered, audit = filter_records(report, docs, model or ModelRouter(self.db, job_id, control)
                                             if any(c["applies_to"] != "institution"
                                                    for c in report.plan.get("criteria", [])) else None,
                                             control, progress)
            docs = sorted(filtered, key=lambda d: (d.source_hospital.casefold(), (d.get("name") or "").casefold()))
            per_hospital = report.plan.get("requested_people_per_institution")
            if per_hospital:
                selected = []
                for source in sources:
                    matched = [d for d in docs if d.source_hospital == source.hospital]
                    if len(matched) < per_hospital:
                        report.issues.append(f"{source.hospital}: requested {per_hospital} doctors; {len(matched)} matched")
                    selected.extend(matched[:per_hospital])
                docs = selected
            count = report.plan.get("requested_people")
            if count and len(docs) < count:
                report.issues.append(f"Requested {count} doctors; {len(docs)} matched with available evidence")
            if count:
                docs = docs[:count]
            matching = {d.source_hospital for d in docs}
            requested = report.plan.get("requested_institutions")
            if requested and len(matching) < requested:
                report.issues.append(f"Requested {requested} hospitals; {len(matching)} produced matching doctors "
                                     "after extraction and filtering")
            if not docs:
                report.issues.append("No matching doctors were extracted; source discovery alone does not establish "
                                     "contact coverage")
            sites = []
            for site in extracted["sites"]:
                matched = [d for d in docs if d.source_hospital == site["hospital"]]
                sites.append({**site, "oncologists": sum(d.is_oncology for d in matched), "doctors": len(matched),
                              "with_email": sum(bool(d.get("email")) for d in matched),
                              "with_phone": sum(bool(d.get("phone")) for d in matched),
                              "with_any_contact": sum(any(d.get(f) for f in CONTACT) for d in matched)})
            limited = any(s["status"] in {"failed", "error", "partial"} for s in sites) or bool(report.issues)
            result = CrawlResult([u for s in sources for u in s.urls], docs, [PageLog(**p) for p in extracted["pages"]])
            onco = result.oncologists
            result.stats = {"prompt": body.prompt, "hospitals": len(sources), "doctors_found": len(docs),
                            "oncologists": len(onco),
                            "oncologists_with_doctor_email": sum(1 for d in onco if d.get("email")),
                            "oncologists_with_doctor_phone": sum(1 for d in onco if d.get("phone")),
                            "oncologists_with_any_contact": sum(1 for d in onco if any(d.get(f) for f in CONTACT)),
                            "hospitals_with_matching_doctors": len(matching), "pages_fetched": len(result.pages),
                            "filtered_out": sum(not a["accepted"] for a in audit),
                            "coverage": "partial / limited" if limited else
                                        "configured run completed; not a guarantee of exhaustive coverage",
                            "discovery_notes": " | ".join(report.issues),
                            "seconds": round(time.time() - job["started"], 1)}
            control()
            self.update(job_id, phase="exporting", progress={})
            path = to_xlsx(result, self.out_dir / f"prompt-{job_id}.xlsx", sites=sites)
            self.db.cache(job_id, "filter_audit", audit)
            self.history.save_doctors(job_id, docs)
            for site in sites:
                self.history.save_site(job_id, site)
            self.update(job_id, state="partial" if limited else "done", phase="finished", xlsx=str(path),
                        report=asdict(report), stats=result.stats, finished=time.time())
        except ServiceWait as exc:
            if self.stop.is_set() or self.closed.is_set():
                state = self.db.get(job_id)["state"]
                self.update(job_id, state="stopped" if state == "stopping" else
                            ("interrupted" if self.closed.is_set() else "paused"), retry_at=None)
            else:
                self.update(job_id, state="waiting", error=str(exc), retry_at=exc.retry_at)
        except Paused:
            state = self.db.get(job_id)["state"]
            self.update(job_id, state="stopped" if state == "stopping" else
                        ("interrupted" if self.closed.is_set() else "paused"), retry_at=None)
        except ConfigurationError as exc:
            fail(str(exc))
        except Exception as exc:
            # Do not expose arbitrary HTTP exception bodies or API credentials.
            fail(f"{type(exc).__name__}: the run could not finish. Progress is saved; resume after checking setup")
        finally:
            self.update(job_id, usage=self.db.stats(job_id))
            for service in (model, search, fetcher):
                if service:
                    service.close()


def export_sources(job, path):
    """The same upload columns, so the found hospitals can be re-run as a normal batch."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from .sources import TEMPLATE_COLUMNS
    report = Report(**job["report"])
    if job["body"].get("engine") == "sites" and report.plan.get("hospitals"):
        hospitals = [Hospital(**h) for h in report.plan["hospitals"]]
        sources = sources_from_hospitals(hospitals)
        homes = {h.name: h.homepage for h in hospitals}
    else:
        sources, homes = sources_from_report(report), {}
    requires_filter = any(c["applies_to"] != "institution" for c in report.plan.get("criteria", []))
    keys = [c[0] for c in TEMPLATE_COLUMNS]
    wb = Workbook()
    ws = wb.active
    ws.title = "Hospitals"
    ws.append([c[1] for c in TEMPLATE_COLUMNS])
    for source in sources:
        row = {"hospital_name": source.hospital, "official_website": homes.get(source.hospital, ""),
               "doctors_list_urls": "\n".join(source.urls), "city": source.city,
               "enabled": "no" if requires_filter else "yes", "max_pages": job["body"]["pages_per_hospital"],
               "notes": ("Use the prompt run to apply doctor-level filters; this row is disabled to prevent an "
                         "unfiltered crawl. " if requires_filter else "") + job["body"]["prompt"]}
        ws.append([str(row.get(k, "") or "") for k in keys])
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="123A44")
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 36
        for cell in col:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return Path(path)
