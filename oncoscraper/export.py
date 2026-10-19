"""Excel / CSV / JSON output."""
from __future__ import annotations

import csv
import io
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .crawl import CrawlResult
from .models import ALL_FIELDS, LIST_FIELDS, DoctorRecord

COLUMNS = ["name", "designation", "specialty", "sub_specialties", "department", "qualifications", "experience_years",
           "hospital", "city", "languages", "expertise", "opd_timings", "consultation_fee", "registration", "phone",
           "email", "education", "awards", "about", "photo_url"]
assert set(COLUMNS) == set(ALL_FIELDS)
META_COLUMNS = ["profile_url", "source_url", "found_on", "is_oncology", "oncology_score", "oncology_terms",
                "page_type", "fields_filled"]


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, list):
        value = "; ".join(str(v) for v in value)
    value = str(value)
    if value[:1] in ("=", "+", "-", "@"):
        value = "'" + value  # never let page text become a spreadsheet formula
    return value[:32000]


def row(rec: DoctorRecord) -> list:
    d = rec.to_dict()
    return [_cell(d.get(c)) for c in COLUMNS] + [
        _cell(rec.profile_url), _cell(rec.source_url), _cell(rec.found_on), "YES" if rec.is_oncology else "no",
        round(rec.oncology_score, 2), _cell(rec.oncology_terms), rec.page_type, len(rec.filled_fields()),
    ]


def _sheet(wb, title, headers, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="164E63")
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.columns):
        letter = col[0].column_letter
        width = (widths or {}).get(headers[i]) or min(60, max(10, max(len(str(c.value or "")) for c in col[:200]) + 2))
        ws.column_dimensions[letter].width = width
    return ws


def to_xlsx(result: CrawlResult, path: str | Path) -> Path:
    path = Path(path)
    wb = Workbook()
    wb.remove(wb.active)
    headers = [c.replace("_", " ").title() for c in COLUMNS + META_COLUMNS]
    widths = {"About": 60, "Expertise": 45, "Profile Url": 45, "Source Url": 45, "Found On": 40}
    onco = sorted(result.oncologists, key=lambda d: (d.get("hospital") or "", d.get("name") or ""))
    _sheet(wb, "Oncologists", headers, [row(d) for d in onco], widths)
    _sheet(wb, "All doctors", headers, [row(d) for d in sorted(result.doctors, key=lambda d: d.get("name") or "")], widths)
    ev_rows = []
    for d in result.doctors:
        for e in d.evidence:
            ev_rows.append([d.get("name"), e.field, _cell(e.value), e.layer, _cell(e.snippet), e.source_url])
    _sheet(wb, "Evidence", ["Doctor", "Field", "Value", "Found by", "Page text", "Page URL"], ev_rows,
           {"Page text": 70, "Value": 40, "Page URL": 45})
    _sheet(wb, "Pages", ["URL", "HTTP status", "Page type", "Doctors on page", "Depth", "Note"],
           [[p.url, p.status, p.page_type, p.doctors, p.depth, p.note] for p in result.pages], {"URL": 70})
    summary = [["Start URL(s)", ", ".join(result.start_urls)], ["Run finished", datetime.now().strftime("%Y-%m-%d %H:%M")]]
    summary += [[k.replace("_", " "), json.dumps(v) if isinstance(v, dict) else v] for k, v in result.stats.items()]
    fill = {f: sum(1 for d in onco if d.get(f)) for f in COLUMNS}
    summary.append(["", ""])
    summary.append(["Field coverage (oncologists)", f"{len(onco)} records"])
    summary += [[f, f"{n} ({(100 * n // len(onco)) if onco else 0}%)"] for f, n in fill.items()]
    _sheet(wb, "Summary", ["Item", "Value"], summary, {"Item": 32, "Value": 80})
    wb.move_sheet("Summary", offset=-4)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def to_csv(doctors: list[DoctorRecord]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(COLUMNS + META_COLUMNS)
    for d in doctors:
        w.writerow(row(d))
    return buf.getvalue()


def to_json(result: CrawlResult, oncology_only: bool = True) -> str:
    docs = result.oncologists if oncology_only else result.doctors
    return json.dumps({"start_urls": result.start_urls, "stats": result.stats,
                       "doctors": [d.to_dict() for d in docs]}, ensure_ascii=False, indent=2)


def write(result: CrawlResult, out: str, oncology_only: bool = True) -> Path:
    p = Path(out)
    suffix = p.suffix.lower()
    if suffix == ".xlsx":
        return to_xlsx(result, p)
    p.parent.mkdir(parents=True, exist_ok=True)
    if suffix == ".csv":
        p.write_text(to_csv(result.oncologists if oncology_only else result.doctors), encoding="utf-8-sig")
    elif suffix == ".json":
        p.write_text(to_json(result, oncology_only), encoding="utf-8")
    else:
        raise ValueError("output must end in .xlsx, .csv or .json")
    return p


__all__ = ["to_xlsx", "to_csv", "to_json", "write", "COLUMNS", "LIST_FIELDS"]
