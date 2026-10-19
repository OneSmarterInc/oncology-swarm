from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select, update
from app.models import Doctor, Proposal, Snapshot, Affiliation, Contact, Registration, DoctorSource, Control, utcnow

def safe_cell(value):
    # Never let externally sourced text become an Excel formula.
    if isinstance(value, str):
        value = ''.join(c for c in value if ord(c) >= 32 or c in '\t\n\r')
        if value.lstrip().startswith(('=', '+', '-', '@')) or value.startswith(('\t', '\r', '\n')):
            value = "'" + value
        return value[:32767]
    return value

def export_xlsx(db):
    wb = Workbook()
    wb.remove(wb.active)
    sheets = {
        'Doctors': ['doctor_id', 'name', 'specialty', 'city', 'synthetic', 'approved_at_utc', 'job_id', 'source_type', 'source_id', 'qualifications', 'experience_years', 'sub_specialty', 'languages', 'opd_timings', 'consultation_fee', 'profile_photo_url'],
        'Affiliations': ['doctor_id', 'institution', 'city', 'evidence_ids', 'hospital_id'],
        'Contacts': ['doctor_id', 'type', 'value', 'scope', 'evidence_ids'],
        'Registration': ['doctor_id', 'number', 'authority', 'status', 'evidence_ids'],
        'Evidence': ['doctor_id', 'job_id', 'evidence_id', 'field', 'value', 'source_url', 'support_text', 'retrieved_at'],
        'Sources': ['doctor_id', 'source_url', 'content_sha256', 'retrieved_at', 'synthetic', 'source_type', 'source_id', 'job_id'],
        'Review': ['doctor_id', 'job_id', 'reviewer', 'reviewed_at_utc', 'note', 'possible_duplicate_job_ids'],
        'Readme': ['topic', 'note'],
    }
    for name, headers in sheets.items():
        wb.create_sheet(name).append(headers)
    def add(sheet, row): wb[sheet].append([safe_cell(v) for v in row])
    add('Readme', ['Scope', 'Only human-approved records are exported. Source evidence is provenance, not a guarantee of truth.'])
    add('Readme', ['Synthetic', 'synthetic=TRUE identifies demonstration records and must never be used for outreach.'])
    add('Readme', ['Contact scope', 'doctor means explicitly labeled professional direct contact; institution_central is a shared hospital/institution channel; unclassified is unresolved.'])
    add('Readme', ['Registration', 'provider_claimed is not an independently verified current license. This starter does not perform registry verification.'])
    add('Readme', ['Duplicates', 'Possible duplicates are flagged. Explicit merge retains all contributing source evidence; target identity fields remain primary. Withdrawn and changed-pending-review records are excluded.'])
    # Share the review lock while materializing the export dataset. This prevents
    # mixing a prior approved identity with a concurrent unreviewed correction.
    # ZIP serialization occurs after releasing the transaction.
    with db.session.begin() as s:
        s.execute(update(Control).where(Control.id==1).values(updated_at=utcnow()))
        def refs(evidence,proposal_id):
            job=s.scalar(select(Proposal.job_id).where(Proposal.id==proposal_id))
            return ', '.join(f"job:{job}:{e['id']}" for e in evidence)
        for d in s.scalars(select(Doctor).where(Doctor.active.is_(True), Doctor.needs_review.is_(False)).order_by(Doctor.id)):
            p = s.get(Proposal, d.proposal_id)
            source = s.scalar(select(Snapshot).where(Snapshot.job_id == p.job_id))
            add('Doctors', [d.id, d.name, d.specialty, d.city, d.synthetic, d.approved_at.isoformat() + 'Z', p.job_id, d.source_type, d.source_id, d.qualifications, d.experience_years, d.sub_specialty, d.languages, d.opd_timings, d.consultation_fee, d.profile_photo_url])
            for a in s.scalars(select(Affiliation).where(Affiliation.doctor_id == d.id)):
                add('Affiliations', [d.id, a.name, a.city, refs(a.evidence,a.source_proposal_id), a.hospital_id])
            for c in s.scalars(select(Contact).where(Contact.doctor_id == d.id)):
                add('Contacts', [d.id, c.type, c.value, c.scope, refs(c.evidence,c.source_proposal_id)])
            for r in s.scalars(select(Registration).where(Registration.doctor_id == d.id)):
                add('Registration', [d.id, r.number, r.authority, r.status, refs(r.evidence,r.source_proposal_id)])
            proposals=s.scalars(select(Proposal).join(DoctorSource, DoctorSource.proposal_id==Proposal.id).where(DoctorSource.doctor_id==d.id).order_by(Proposal.id)).all()
            for contributing in proposals or [p]:
                source=s.scalar(select(Snapshot).where(Snapshot.job_id==contributing.job_id))
                for e in contributing.data['evidence']:
                    add('Evidence', [d.id, contributing.job_id, f"job:{contributing.job_id}:{e['id']}", e['field'], e['value'], e['source_url'], e['support_text'], e['retrieved_at']])
                add('Sources', [d.id, source.url, source.content_hash, source.retrieved_at, source.synthetic, source.source_type, source.source_id, contributing.job_id])
                add('Review', [d.id, contributing.job_id, contributing.reviewer, contributing.reviewed_at.isoformat()+'Z', contributing.review_note, ', '.join(map(str, contributing.duplicate_ids))])
    for ws in wb:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='164E63')
        for column in ws.columns:
            width = min(65, max(14, max(len(str(c.value or '')) for c in column) + 2))
            ws.column_dimensions[column[0].column_letter].width = width
    result = BytesIO(); wb.save(result)
    return result.getvalue()
